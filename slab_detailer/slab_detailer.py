#!/usr/bin/env python3
"""
slab_detailer.py - Flat slab reinforcement DETAILING to Eurocode 2 (EN 1992-1-1)

The program does NOT design the slab. You give the reinforcement in config.yaml
(meshes, bands, extra bars over columns, ...). The program reads the slab geometry
from your DXF and:

  * lays every family of bars in both directions, clipped to the slab edges and openings
  * cuts long bars to the stock length and laps them where EC2 good practice wants it
      - bottom bars: laps over the support lines (columns / walls)
      - top bars:    laps at mid-span
      - laps staggered (alternate bars), lap length from EC2 8.4 / 8.7
  * places the extra top bars over columns between the band bars (interleaved)
  * classifies columns as internal / edge / corner automatically
  * adds integrity bars through columns, trimmers at openings, U-bars at free edges
  * draws the punching control perimeter u1 at 2d
  * writes: all-layers DXF, TOP and BOTTOM DXFs, bar schedule (xlsx), report (txt),
    PNG previews

Mat foundation: mat_foundation.py runs this same engine with structure: mat_foundation
(config_mat.yaml) - top bars lapped over the supports, bottom bars at mid-span.

Usage:
    python slab_detailer.py my_slab.dxf                      (uses config.yaml next to it)
    python slab_detailer.py my_slab.dxf -c my_config.yaml -o output_folder
"""
from __future__ import annotations

import argparse
import math
import os
import sys
from collections import defaultdict
from dataclasses import dataclass, field

import ezdxf
import yaml
from ezdxf import path as ezpath
from shapely.geometry import LineString, MultiPolygon, Point, Polygon, box
from ezdxf.render.arrows import ARROWS
from shapely.ops import polygonize, unary_union

UNIT_MM = {"mm": 1.0, "cm": 10.0, "m": 1000.0}
WARNINGS: list[str] = []
STRUCTURE = ["flat_slab"]          # flat_slab | mat_foundation (set by run() from the config)


def is_mat() -> bool:
    """Mat foundation: an upside-down flat slab (soil pressure up, supports down). Tension is at the
    bottom over the supports and at the top in the spans, so compared with a flat slab the laps swap
    (top bars lapped over the supports, bottom bars at mid-span) and the extra bars at the supports
    are bottom bars."""
    return STRUCTURE[0] == "mat_foundation"


def warn(msg: str):
    if msg not in WARNINGS:
        WARNINGS.append(msg)
        print("  WARNING:", msg)


def ceil50(v: float) -> float:
    return math.ceil(v / 50.0 - 1e-9) * 50.0


# =============================================================================
# Parameters
# =============================================================================
class Params:
    def __init__(self, cfg: dict):
        s = cfg["slab"]
        m = cfg["materials"]
        self.cfg = cfg
        self.h = float(s["thickness"])
        self.cb = float(s["cover_bottom"])
        self.ct = float(s["cover_top"])
        self.cs = float(s.get("cover_side", 30))
        d = s.get("effective_depth", "auto")
        self.d = self.h - self.ct - 12.0 if d in (None, "auto") else float(d)
        self.fck = float(m["fck"])
        self.fyk = float(m["fyk"])
        self.gc = float(m.get("gamma_c", 1.5))
        self.gs = float(m.get("gamma_s", 1.15))
        self.fyd = self.fyk / self.gs
        self.fctm = 0.3 * self.fck ** (2 / 3) if self.fck <= 50 else 2.12 * math.log(1 + (self.fck + 8) / 10)
        self.fctk005 = 0.7 * self.fctm
        self.fctd = 1.0 * self.fctk005 / self.gc
        self.stock = float(cfg.get("stock_length", 12000))
        self.stagger = bool(cfg.get("stagger_laps", True))
        self.hooks = bool(cfg.get("top_bar_hooks_at_free_edges", True))
        self.leg = self.h - self.ct - self.cb          # vertical leg of a bend-down
        self.lap_override = cfg.get("lap_lengths")
        if not isinstance(self.lap_override, dict):
            self.lap_override = {}

    # ---- EC2 8.4 / 8.7 -------------------------------------------------------
    def bond(self, dia: float, face: str):
        # Fig. 8.2: in slabs with h > 250 mm, bars more than 250 mm above the
        # bottom are in "poor" bond conditions (eta1 = 0.7)
        eta1 = 0.7 if (face == "T" and self.h > 250 and (self.h - self.ct) > 250) else 1.0
        eta2 = 1.0 if dia <= 32 else (132 - dia) / 100
        fbd = 2.25 * eta1 * eta2 * self.fctd
        lb_rqd = dia / 4 * self.fyd / fbd
        return eta1, fbd, lb_rqd

    def lbd(self, dia: float, face: str) -> float:
        _, _, lb_rqd = self.bond(dia, face)
        return ceil50(max(lb_rqd, 0.3 * lb_rqd, 10 * dia, 100))

    def lap(self, dia: float, face: str) -> float:
        if int(dia) in self.lap_override:
            return float(self.lap_override[int(dia)])
        _, _, lb_rqd = self.bond(dia, face)
        a6 = 1.4 if self.stagger else 1.5          # Table 8.3: <=50% / >50% lapped
        l0 = max(a6 * lb_rqd, 0.3 * a6 * lb_rqd, 15 * dia, 200)
        return ceil50(l0)


# =============================================================================
# Reading the DXF
# =============================================================================
def _expand(e, parent_layer):
    lay = e.dxf.layer if e.dxf.hasattr("layer") else parent_layer
    if lay == "0":
        lay = parent_layer
    if e.dxftype() == "INSERT":
        try:
            for v in e.virtual_entities():
                yield from _expand(v, lay)
        except Exception:
            return
    else:
        yield e, lay


def iter_entities(layout):
    for e in layout:
        yield from _expand(e, e.dxf.layer)


def mesh_faces_on_layers(ents, names, u):
    """3DFACE / SOLID / TRACE elements (e.g. a slab exported from SAP2000 / ETABS as a FE mesh)."""
    names = {n.upper() for n in names}
    faces = []
    for e, lay in ents:
        if lay.upper() not in names or e.dxftype() not in ("3DFACE", "SOLID", "TRACE"):
            continue
        vs = [e.dxf.get(f"vtx{i}") for i in range(4)]
        if e.dxftype() != "3DFACE":                 # SOLID/TRACE vertex order is 0-1-3-2
            vs = [vs[0], vs[1], vs[3], vs[2]]
        pts = []
        for v in vs:
            if v is None:
                continue
            q = (v.x * u, v.y * u)
            if not pts or math.dist(q, pts[-1]) > 0.01:
                pts.append(q)
        if len(pts) > 2 and math.dist(pts[0], pts[-1]) < 0.01:
            pts.pop()
        if len(pts) >= 3:
            pg = Polygon(pts).buffer(0)
            if pg.area > 1:
                faces.append(pg)
    return faces


def shapes_on_layers(ents, names, u):
    names = {n.upper() for n in names}
    polys, lines = [], []
    for e, lay in ents:
        if lay.upper() not in names:
            continue
        t = e.dxftype()
        if t == "HATCH":                       # filled columns / walls
            try:
                for hp in ezpath.from_hatch(e):
                    pts = [(v.x * u, v.y * u) for v in hp.flattening(distance=0.5 / u)]
                    if len(pts) >= 3:
                        pg = Polygon(pts).buffer(0)
                        polys += [g for g in getattr(pg, "geoms", [pg]) if g.area > 100]
            except Exception as ex:
                warn(f"could not read a HATCH on layer {lay}: {ex}")
            continue
        if t not in ("LINE", "ARC", "LWPOLYLINE", "POLYLINE", "CIRCLE", "ELLIPSE", "SPLINE"):
            continue
        try:
            p = ezpath.make_path(e)
            pts = [(v.x * u, v.y * u) for v in p.flattening(distance=0.5 / u)]
        except Exception as ex:
            warn(f"could not read a {t} on layer {lay}: {ex}")
            continue
        if len(pts) < 2:
            continue
        closed = False
        if t == "CIRCLE":
            closed = True
        elif t == "LWPOLYLINE":
            closed = e.closed
        elif t == "POLYLINE":
            closed = e.is_closed
        elif t == "ELLIPSE":
            closed = abs(abs(e.dxf.end_param - e.dxf.start_param) - 2 * math.pi) < 1e-6
        if t not in ("LINE", "ARC") and math.dist(pts[0], pts[-1]) < 1.0:
            closed = True
        if closed and len(pts) >= 3:
            pg = Polygon(pts).buffer(0)
            for g in getattr(pg, "geoms", [pg]):
                if g.area > 100:
                    polys.append(g)
        else:
            lines.append(LineString(pts))
    if lines:
        for pg in polygonize(unary_union(lines)):
            if pg.area > 100:
                polys.append(pg)
    return polys


def close_along_walls(ents, slab_names, wall_names, u, tol=10.0):
    """A slab outline drawn the usual way, stopping at the faces of the walls and columns: the slab
    lines and the outlines of the walls/columns (wall_names: their layers) together close it (gaps up
    to 2 x tol are bridged). The walls and columns are part of the slab."""
    names = {n.upper() for n in slab_names + wall_names}
    lines = []
    for e, lay in ents:
        if lay.upper() not in names or e.dxftype() not in ("LINE", "ARC", "LWPOLYLINE", "POLYLINE", "SPLINE"):
            continue
        try:
            pts = [(v.x * u, v.y * u) for v in ezpath.make_path(e).flattening(distance=0.5 / u)]
        except Exception:
            continue
        if len(pts) >= 2:
            lines.append(LineString(pts))
    if not lines:
        return []
    walls = shapes_on_layers(ents, wall_names, u)
    walls_u = nest(walls) if walls else Polygon()
    blob = unary_union([ln.buffer(tol, join_style=2) for ln in lines])        # round ends bridge the gaps
    regions = []
    for g in getattr(blob, "geoms", [blob]):
        for r in g.interiors:
            pg = Polygon(r)
            # the inside of a wall outline is not slab (free-standing walls are islands in a region)
            if pg.area > 0.5e6 and pg.intersection(walls_u).area < 0.5 * pg.area:
                regions.append(pg.buffer(tol, join_style=2))
    if not regions:
        return []
    near = [w for w in getattr(walls_u, "geoms", [walls_u]) if not w.is_empty and
            any(w.distance(r) < 2 * tol for r in regions)]
    slab = unary_union(regions + near).buffer(2 * tol, join_style=2).buffer(-2 * tol, join_style=2)
    return [g for g in getattr(slab, "geoms", [slab]) if g.area > 0.5e6]


def points_to_supports(ents, names, u, spec, slab):
    """Support nodes of an FE model (POINT entities): points in a row with a spacing up to
    'max_spacing' form a wall (extended by half a spacing at each end), single points are columns.
    Wall thickness and column size are not in the file, so they come from the config."""
    if not names:
        return [], []
    names_u = {n.upper() for n in names}
    pts = []
    for e, lay in ents:
        if lay.upper() in names_u and e.dxftype() == "POINT":
            p = e.dxf.location
            pts.append((round(p.x * u, 1), round(p.y * u, 1)))
    pts = sorted(set(pts))
    if not pts:
        return [], []
    t = float(spec.get("wall_thickness", 250))
    cs = spec.get("column_size", [400, 400])
    cs = [float(cs), float(cs)] if not isinstance(cs, (list, tuple)) else [float(cs[0]), float(cs[-1])]
    smax = float(spec.get("max_spacing", 700))
    tol = 20.0
    in_wall = set()
    walls = []
    for axis in (1, 0):                       # rows along X (same y), then along Y (same x)
        lines = defaultdict(list)
        for p in pts:
            lines[round(p[axis] / tol)].append(p)
        for key, lp in lines.items():
            lp.sort(key=lambda q: q[1 - axis])
            run = [lp[0]]
            for a, b in zip(lp, lp[1:]):
                if b[1 - axis] - a[1 - axis] <= smax:
                    run.append(b)
                else:
                    if len(run) >= 2:
                        walls.append((axis, run))
                    run = [b]
            if len(run) >= 2:
                walls.append((axis, run))
    polys = []
    for axis, run in walls:
        c = sum(p[axis] for p in run) / len(run)
        lo, hi = run[0][1 - axis], run[-1][1 - axis]
        sp = (hi - lo) / (len(run) - 1)
        lo, hi = lo - sp / 2, hi + sp / 2
        r = box(lo, c - t / 2, hi, c + t / 2) if axis == 1 else box(c - t / 2, lo, c + t / 2, hi)
        r = r.intersection(slab.buffer(t)) if not slab.is_empty else r
        if r.area > 0:
            polys.append(r)
            in_wall.update(run)
    cols = []
    wall_u = unary_union(polys) if polys else Polygon()
    for p in pts:
        if p in in_wall or (not wall_u.is_empty and wall_u.buffer(10).contains(Point(p))):
            continue
        cols.append(box(p[0] - cs[0] / 2, p[1] - cs[1] / 2, p[0] + cs[0] / 2, p[1] + cs[1] / 2))
    return polys, cols


def nest(polys):
    """Even-odd nesting: an outline inside another one is a hole of it."""
    polys = sorted(polys, key=lambda p: p.area, reverse=True)
    depth = []
    for i, p in enumerate(polys):
        rp = p.representative_point()
        d = sum(1 for j in range(i) if polys[j].contains(rp))
        depth.append(d)
    solid = unary_union([p for p, d in zip(polys, depth) if d % 2 == 0])
    holes = unary_union([p for p, d in zip(polys, depth) if d % 2 == 1])
    return solid.difference(holes) if not holes.is_empty else solid


@dataclass
class Column:
    id: int
    poly: Polygon
    cx: float
    cy: float
    wx: float
    wy: float
    kind: str = "internal"
    round: bool = False
    near_opening: bool = False


@dataclass
class WallRect:
    id: int
    poly: Polygon
    orient: str | None      # 'x' = long side along X, 'y' = along Y
    cx: float
    cy: float
    length: float
    thk: float
    lo: float               # extent along the long axis
    hi: float
    core: bool = False      # piece of a closed box of walls (core)


def rectilinear_pieces(poly: Polygon):
    """Split an orthogonal wall outline (L, U, box with hole...) into straight rectangles."""
    coords = list(poly.exterior.coords)
    for r in poly.interiors:
        coords += list(r.coords)
    # corners drawn a fraction of a mm apart are one line: otherwise they give zero-thickness pieces
    xs = cluster([x for x, _ in coords], 5.0)
    ys = cluster([y for _, y in coords], 5.0)
    inside = {}
    for i in range(len(xs) - 1):
        for j in range(len(ys) - 1):
            c = Point((xs[i] + xs[i + 1]) / 2, (ys[j] + ys[j + 1]) / 2)
            inside[(i, j)] = poly.contains(c)
    used = set()
    rects = []
    # long horizontal runs first, then vertical runs from the remaining cells
    for j in range(len(ys) - 1):
        i = 0
        while i < len(xs) - 1:
            if inside[(i, j)] and (i, j) not in used:
                k = i
                while k + 1 < len(xs) - 1 and inside[(k + 1, j)]:
                    k += 1
                if xs[k + 1] - xs[i] >= 2 * (ys[j + 1] - ys[j]):
                    rects.append(box(xs[i], ys[j], xs[k + 1], ys[j + 1]))
                    used.update((q, j) for q in range(i, k + 1))
                i = k + 1
            else:
                i += 1
    for i in range(len(xs) - 1):
        j = 0
        while j < len(ys) - 1:
            if inside[(i, j)] and (i, j) not in used:
                k = j
                while k + 1 < len(ys) - 1 and inside[(i, k + 1)] and (i, k + 1) not in used:
                    k += 1
                rects.append(box(xs[i], ys[j], xs[i + 1], ys[k + 1]))
                used.update((i, q) for q in range(j, k + 1))
                j = k + 1
            else:
                j += 1
    return [r for r in rects if min(r.bounds[2] - r.bounds[0], r.bounds[3] - r.bounds[1]) >= 20]


def is_orthogonal(poly: Polygon) -> bool:
    rings = [poly.exterior] + list(poly.interiors)
    for r in rings:
        c = list(r.coords)
        for (x0, y0), (x1, y1) in zip(c, c[1:]):
            if abs(x1 - x0) > 1 and abs(y1 - y0) > 1:
                return False
    return True


def make_wallrect(i, r: Polygon) -> WallRect:
    mrr = r.minimum_rotated_rectangle
    c = list(mrr.exterior.coords)
    e1 = (c[1][0] - c[0][0], c[1][1] - c[0][1])
    e2 = (c[2][0] - c[1][0], c[2][1] - c[1][1])
    lg, sh = (e1, e2) if math.hypot(*e1) >= math.hypot(*e2) else (e2, e1)
    ang = math.degrees(math.atan2(lg[1], lg[0])) % 180
    orient = "x" if (ang < 2 or ang > 178) else ("y" if abs(ang - 90) < 2 else None)
    minx, miny, maxx, maxy = r.bounds
    ct = r.centroid
    if orient == "x":
        lo, hi = minx, maxx
    elif orient == "y":
        lo, hi = miny, maxy
    else:
        lo = hi = 0
    return WallRect(i, r, orient, ct.x, ct.y, math.hypot(*lg), math.hypot(*sh), lo, hi)


class Geometry:
    def __init__(self, doc, cfg, P: Params):
        u = UNIT_MM[cfg.get("units", "mm")]
        self.u = u
        L = cfg["layers"]

        def lay(k):
            v = L.get(k) or []
            return [v] if isinstance(v, str) else list(v)

        ents = list(iter_entities(doc.modelspace()))

        # ---- slab -----------------------------------------------------------
        self.from_mesh = False
        faces = mesh_faces_on_layers(ents, lay("slab"), u)
        sp = shapes_on_layers(ents, lay("slab"), u)
        if faces:
            # FE mesh: the slab is the union of all elements; gaps between them are openings
            self.from_mesh = True
            mesh = unary_union([f.buffer(1, join_style=2) for f in faces]).buffer(-1, join_style=2)
            cleaned = []
            for g in getattr(mesh, "geoms", [mesh]):
                holes = [r for r in g.interiors if Polygon(r).area > 0.05e6]      # ignore slivers < 0.05 m2
                cleaned.append(Polygon(g.exterior, holes).simplify(5))
            mesh = unary_union(cleaned)
            sp = sp + [mesh]
            print(f"  slab read from {len(faces)} mesh elements")
        if not sp:
            sp = close_along_walls(ents, lay("slab"), lay("walls") + lay("columns"), u)
            if sp:
                print("  slab outline closed along the wall faces")
        if not sp:
            sys.exit(f"ERROR: no closed slab outline found on layer(s) {lay('slab')}. "
                     f"Check 'layers: slab' in the config.")
        slab_with_holes = nest(sp) if not faces else unary_union(sp)
        op = shapes_on_layers(ents, lay("openings"), u)
        openings = unary_union(op) if op else Polygon()
        slab = slab_with_holes.difference(openings) if not openings.is_empty else slab_with_holes
        if isinstance(slab, MultiPolygon):
            warn(f"the slab consists of {len(slab.geoms)} separate parts - all are detailed")
        self.slab = slab
        self.inner = slab.buffer(-P.cs, join_style=2)          # where bars may lie
        self.holes = []
        for g in getattr(slab, "geoms", [slab]):
            self.holes += [Polygon(r) for r in g.interiors]

        # ---- supports given as points (FE model export) ------------------------
        self.from_points = False
        pw, pc = points_to_supports(ents, lay("support_points"), u, cfg.get("support_points") or {}, slab)
        if pw or pc:
            self.from_points = True
            print(f"  supports from points: {len(pw)} wall(s), {len(pc)} column(s)")

        # ---- walls ----------------------------------------------------------
        wp = shapes_on_layers(ents, lay("walls"), u) + pw
        walls_union = nest(wp) if wp else Polygon()
        self.walls_union = walls_union
        self.walls: list[WallRect] = []
        for g in getattr(walls_union, "geoms", [walls_union]):
            if g.is_empty:
                continue
            pieces = rectilinear_pieces(g) if is_orthogonal(g) else [g]
            for r in pieces:
                w = make_wallrect(len(self.walls) + 1, r)
                w.core = len(g.interiors) > 0
                if w.orient is None:
                    warn(f"wall near ({w.cx:.0f},{w.cy:.0f}) is not parallel to X or Y - no wall bars")
                self.walls.append(w)

        # ---- columns --------------------------------------------------------
        cp = shapes_on_layers(ents, lay("columns"), u) + pc
        self.columns: list[Column] = []
        skipped = 0
        for p in cp:
            if not walls_union.is_empty and p.intersection(walls_union).area > 0.5 * p.area:
                skipped += 1
                continue
            minx, miny, maxx, maxy = p.bounds
            c = p.centroid
            rnd = len(p.exterior.coords) > 12 and abs(p.area - math.pi * ((maxx - minx) / 2) ** 2) < 0.05 * p.area
            self.columns.append(Column(len(self.columns) + 1, p, c.x, c.y, maxx - minx, maxy - miny, round=rnd))
        if skipped:
            warn(f"{skipped} column(s) lie inside walls and were ignored")
        if not self.columns:
            warn("no columns found - check 'layers: columns'")

        self.classify_columns(cfg, P)

        # ---- support lines (for laps, spans and bands) ----------------------
        tol = 300.0
        self.col_lines = {"x": cluster([c.cx for c in self.columns], tol),
                          "y": cluster([c.cy for c in self.columns], tol)}
        wl_x = [w.cx for w in self.walls if w.orient == "y"]
        wl_y = [w.cy for w in self.walls if w.orient == "x"]
        self.sup_lines = {"x": cluster([c.cx for c in self.columns] + wl_x, tol),
                          "y": cluster([c.cy for c in self.columns] + wl_y, tol)}
        minx, miny, maxx, maxy = self.slab.bounds
        self.bounds = (minx, miny, maxx, maxy)
        self.axes = read_axes(ents, lay("axes") or ["AXIS", "AXIS LINES", "AXES", "GRID"], u, self.bounds)

    def classify_columns(self, cfg, P):
        tol = float(cfg.get("top_column_bars", {}).get("edge_distance", 500))
        outer = unary_union([Polygon(g.exterior) for g in getattr(self.slab, "geoms", [self.slab])])
        edge = outer.boundary
        for c in self.columns:
            near = c.poly.buffer(tol).intersection(edge)
            dirs = set()
            for g in getattr(near, "geoms", [near]):
                if g.is_empty or g.geom_type != "LineString":
                    continue
                cs = list(g.coords)
                for (x0, y0), (x1, y1) in zip(cs, cs[1:]):
                    if math.hypot(x1 - x0, y1 - y0) < 50:
                        continue
                    a = math.degrees(math.atan2(y1 - y0, x1 - x0)) % 180
                    dirs.add(round(a / 30) % 6)
            groups = []
            for d in sorted(dirs):
                if not any(min(abs(d - g), 6 - abs(d - g)) <= 1 for g in groups):
                    groups.append(d)
            c.kind = "internal" if not groups else ("edge" if len(groups) == 1 else "corner")
            # EC2 6.4.2(3): opening closer than 6d to the column reduces the control perimeter
            for k, hole in enumerate(self.holes):
                dist = c.poly.distance(hole)
                if dist <= 6 * P.d:
                    c.near_opening = True
                    warn(f"column {c.id} at ({c.cx:.0f},{c.cy:.0f}) is {dist:.0f} mm from opening {k + 1} "
                         f"(< 6d = {6 * P.d:.0f} mm): reduce u1 for punching (EC2 6.4.2(3), Fig. 6.14)")


def read_axes(ents, names, u, bounds):
    """Grid lines on the axis layer(s): {'x': [(x, ymin, ymax)], 'y': [(y, xmin, xmax)]} in mm."""
    names = {n.upper() for n in names}
    minx, miny, maxx, maxy = bounds
    W, H = maxx - minx, maxy - miny
    vx, hy = [], []
    for e, lay in ents:
        if lay.upper() not in names or e.dxftype() not in ("LINE", "LWPOLYLINE", "POLYLINE"):
            continue
        try:
            pts = [(v.x * u, v.y * u) for v in ezpath.make_path(e).flattening(distance=1)]
        except Exception:
            continue
        for (x0, y0), (x1, y1) in zip(pts, pts[1:]):
            if abs(x1 - x0) < 1 and abs(y1 - y0) > 0.25 * H:
                vx.append(((x0 + x1) / 2, min(y0, y1), max(y0, y1)))
            elif abs(y1 - y0) < 1 and abs(x1 - x0) > 0.25 * W:
                hy.append(((y0 + y1) / 2, min(x0, x1), max(x0, x1)))

    def merge(items):
        items.sort()
        out = []
        for c, lo, hi in items:
            if out and abs(c - out[-1][0]) < 50:
                out[-1] = (out[-1][0], min(lo, out[-1][1]), max(hi, out[-1][2]))
            else:
                out.append((c, lo, hi))
        return out
    return {"x": merge(vx), "y": merge(hy)}


def cluster(vals, tol):
    vals = sorted(vals)
    out, grp = [], []
    for v in vals:
        if grp and v - grp[-1] > tol:
            out.append(sum(grp) / len(grp))
            grp = []
        grp.append(v)
    if grp:
        out.append(sum(grp) / len(grp))
    return out


def neighbours(lines, c, tol=300.0):
    prev = [v for v in lines if v < c - tol]
    nxt = [v for v in lines if v > c + tol]
    return (max(prev) if prev else None), (min(nxt) if nxt else None)


# =============================================================================
# Bars
# =============================================================================
@dataclass
class Bar:
    fam: str            # MESH, BAND, COL, WALL, INTEG, TRIM
    face: str           # 'B' bottom / 'T' top
    dirn: str           # 'x' or 'y' = direction the bar runs in
    across: float       # coordinate perpendicular to the bar
    a: float            # start (along)
    b: float            # end (along)
    dia: int
    s: float            # nominal spacing
    tag: str = ""
    hook_a: float = 0.0
    hook_b: float = 0.0
    chain: int = -1     # id of the continuous run this bar belongs to (lapped bars share it)
    ci: int = 0         # position in that run

    @property
    def length(self):
        return (self.b - self.a) + self.hook_a + self.hook_b


def pt(dirn, along, across):
    return (along, across) if dirn == "x" else (across, along)


def clip_segments(region, dirn, across, a, b, minlen=50.0):
    ln = LineString([pt(dirn, a, across), pt(dirn, b, across)])
    g = region.intersection(ln)
    out = []

    def walk(geom):
        if geom.is_empty:
            return
        if geom.geom_type == "LineString":
            if geom.length >= minlen:
                cs = [c[0] if dirn == "x" else c[1] for c in geom.coords]
                # snap the bar ends inwards to whole cm: a slightly skewed edge then does not give
                # every bar a different length (and a different bar mark)
                lo_, hi_ = math.ceil(min(cs) / 10 - 1e-6) * 10, math.floor(max(cs) / 10 + 1e-6) * 10
                if hi_ - lo_ >= minlen:
                    out.append((lo_, hi_))
        elif hasattr(geom, "geoms"):
            for gg in geom.geoms:
                walk(gg)

    walk(g)
    out.sort()
    # merge touching pieces
    merged = []
    for s in out:
        if merged and s[0] <= merged[-1][1] + 1:
            merged[-1] = (merged[-1][0], max(merged[-1][1], s[1]))
        else:
            merged.append(s)
    return merged


def lap_zones(lines, l0, f, midspan=False):
    """Allowed ranges for the CENTRE of a lap, so that the whole lap lies in the low-stress zone:
    bottom bars: from f x (left span) before a support line to f x (right span) after it,
    top bars:    +- f x span around mid-span."""
    zones = []
    if midspan:
        for p, q in zip(lines, lines[1:]):
            m, z = (p + q) / 2, f * (q - p)
            zones.append((m - z + l0 / 2, m + z - l0 / 2) if z > l0 / 2 else (m, m))
    else:
        for i, c in enumerate(lines):
            zl = f * (c - lines[i - 1]) if i > 0 else 0.0
            zr = f * (lines[i + 1] - c) if i + 1 < len(lines) else 0.0
            lo, hi = c - zl + l0 / 2, c + zr - l0 / 2
            zones.append((lo, hi) if hi >= lo else (c, c))
    return zones


def split_bar(a, b, zones, stock, l0, shift, label):
    """Cut a run [a,b] into bars <= stock, each lap placed as far along as possible inside a lap zone."""
    pieces = []
    s = a
    guard = 0
    while b - s > stock + 1e-6 and guard < 200:
        guard += 1
        best = None
        for lim in ((s + stock - l0 / 2 - shift), (s + stock - l0 / 2)) if shift else ((s + stock - l0 / 2),):
            for lo, hi in zones:
                c = min(hi, lim)
                if c >= max(lo, s + l0) and (b - (c - l0 / 2)) >= l0 + 300:   # last bar >= l0 + 30 cm
                    best = c if best is None else max(best, c)
            if best is not None:
                break
        if best is None:
            best = s + stock - l0 / 2
            warn(f"{label}: no lap zone within one bar length ({stock:.0f} mm) - some laps placed outside "
                 f"the preferred zones; check them on the plan")
        c = math.floor(best / 10) * 10
        pieces.append((s, c + l0 / 2))
        s = c - l0 / 2
    pieces.append((s, b))
    return pieces


def across_range(G: Geometry, dirn):
    minx, miny, maxx, maxy = G.inner.bounds
    return (miny, maxy) if dirn == "x" else (minx, maxx)


def along_range(G: Geometry, dirn):
    minx, miny, maxx, maxy = G.inner.bounds
    return (minx, maxx) if dirn == "x" else (miny, maxy)


def mesh_positions(G, dirn, s):
    lo, hi = across_range(G, dirn)
    n = int(math.floor((hi - lo) / s + 1e-6))
    return [lo + i * s for i in range(n + 1)]


_CHAIN = [0]


def gen_mesh(G, P, spec, face, dirn, fam, lines, midspan):
    bars = []
    dia, s = int(spec["dia"]), float(spec["spacing"])
    l0 = P.lap(dia, face)
    cands = lap_zones(lines, l0, float(P.cfg.get("lap_zone", 0.2)), midspan)
    shift = ceil50(1.3 * l0) if P.stagger else 0.0
    alo, ahi = along_range(G, dirn)
    stock = P.stock - (2 * P.leg if (face == "T" and P.hooks) else 0)
    for i, pos in enumerate(mesh_positions(G, dirn, s)):
        for (a, b) in clip_segments(G.inner, dirn, pos, alo - 10, ahi + 10,
                                    minlen=max(4 * dia, float(P.cfg.get("min_bar_length", 500)))):
            sh = shift if (i % 2 == 1) else 0.0
            pieces = split_bar(a, b, cands, stock, l0, sh, f"{fam} {face}{dirn}")
            _CHAIN[0] += 1
            for k, (pa, pb) in enumerate(pieces):
                bars.append(Bar(fam, face, dirn, pos, pa, pb, dia, s,
                                chain=_CHAIN[0] if len(pieces) > 1 else -1, ci=k))
    return bars


def ext_len(prev, c, factor, emin):
    if prev is None:
        return 1e7                                  # no support beyond -> run to the slab edge
    return ceil50(max(factor * abs(c - prev), emin))


def top_layout(cfg):
    """mesh | bands_and_distribution | bands"""
    lay = cfg.get("top_layout")
    if lay in ("mesh", "bands_and_distribution", "bands"):
        return lay
    if any((cfg.get("top_mesh") or {}).get(d) for d in ("x", "y")):
        return "mesh"
    return "bands"


def gen_grid_bands(G, P, cfg):
    spec = cfg.get("top_grid_bands") or {}
    if top_layout(cfg) == "mesh" or not spec.get("enabled", True):
        return []
    bars = []
    f = float(spec.get("extension_factor", 0.25))
    emin = float(spec.get("min_extension", 1000))
    for dirn in ("x", "y"):
        ds = spec.get(dirn)
        if not ds:
            continue
        dia, s = int(ds["dia"]), float(ds["spacing"])
        for c in G.col_lines[dirn]:
            prev, nxt = neighbours(G.sup_lines[dirn], c)
            el, er = ext_len(prev, c, f, emin), ext_len(nxt, c, f, emin)
            for pos in mesh_positions(G, dirn, s):
                for (a, b) in clip_segments(G.inner, dirn, pos, c - el, c + er, minlen=4 * dia):
                    if a <= c + 1 and b >= c - 1:
                        bars.append(Bar("BAND", "T", dirn, pos, a, b, dia, s, tag=f"L{c:.0f}"))
    return merge_overlapping_bands(bars, G, P, cfg)


def merge_overlapping_bands(bars, G, P, cfg):
    """Where the bands of two close grid lines overlap on the same bar line, make them one bar
    (no double steel, no meaningless 'laps'); cut it to the stock length if needed."""
    lines = defaultdict(list)
    for b in bars:
        lines[(b.dirn, round(b.across))].append(b)
    out = []
    stock = P.stock - (2 * P.leg if P.hooks else 0)
    zf = float(cfg.get("lap_zone", 0.2))
    for (dirn, _), bl in lines.items():
        bl.sort(key=lambda x: x.a)
        cur = [bl[0]]
        for b in bl[1:]:
            if b.a <= max(x.b for x in cur) + 1:
                cur.append(b)
            else:
                out += _merged(cur, G, P, stock, zf)
                cur = [b]
        out += _merged(cur, G, P, stock, zf)
    return out


def _merged(group, G, P, stock, zf):
    if len(group) == 1:
        return group
    b0 = group[0]
    a, b = min(x.a for x in group), max(x.b for x in group)
    tag = "+".join(sorted({x.tag for x in group}))
    l0 = P.lap(b0.dia, "T")
    zones = lap_zones(G.sup_lines[b0.dirn], l0, zf, midspan=True)
    pieces = split_bar(a, b, zones, stock, l0, 0.0, f"BAND T{b0.dirn}")
    _CHAIN[0] += 1
    cid = _CHAIN[0] if len(pieces) > 1 else -1
    return [Bar("BAND", "T", b0.dirn, b0.across, pa, pb, b0.dia, b0.s, tag=tag, chain=cid, ci=k)
            for k, (pa, pb) in enumerate(pieces)]


def top_base(cfg, dirn):
    """Spacing of the basic bars the extra column bars are placed between (top face; bottom in a mat)."""
    if is_mat():
        bm = (cfg.get("bottom_mesh") or {}).get(dirn)
        return float(bm["spacing"]) if bm else None
    if top_layout(cfg) == "mesh":
        tm = (cfg.get("top_mesh") or {}).get(dirn)
        return float(tm["spacing"]) if tm else None
    gb = cfg.get("top_grid_bands") or {}
    if gb.get("enabled", True) and gb.get(dirn):
        return float(gb[dirn]["spacing"])
    return None


def gen_top_distribution(G, P, cfg, bars_so_far):
    """Distribution bars filling the top face BETWEEN the grid bands. Each one runs l0 into the
    neighbouring band, so it laps with the band bar on the same line (one continuous top layer)."""
    if top_layout(cfg) != "bands_and_distribution":
        return []
    spec = cfg.get("top_distribution_bars") or {}
    out = []
    zf = float(cfg.get("lap_zone", 0.2))
    stock = P.stock - (2 * P.leg if P.hooks else 0)
    for dirn in ("x", "y"):
        ds = spec.get(dirn, spec) if isinstance(spec.get(dirn, spec), dict) else spec
        if not ds or "dia" not in ds:
            continue
        dia, s = int(ds["dia"]), float(ds["spacing"])
        l0 = P.lap(dia, "T")
        zones = lap_zones(G.sup_lines[dirn], l0, zf, midspan=True)
        bands = sorted((b for b in bars_so_far if b.fam == "BAND" and b.dirn == dirn), key=lambda b: b.across)
        bs = float(((cfg.get("top_grid_bands") or {}).get(dirn) or {}).get("spacing", s))
        alo, ahi = along_range(G, dirn)
        import bisect
        acr = [b.across for b in bands]
        for pos in mesh_positions(G, dirn, s):
            i0 = bisect.bisect_left(acr, pos - bs / 2 + 1)
            i1 = bisect.bisect_right(acr, pos + bs / 2 - 1)
            cov = sorted((bands[i].a, bands[i].b) for i in range(i0, i1))
            for (a, b) in clip_segments(G.inner, dirn, pos, alo - 10, ahi + 10, minlen=4 * dia):
                ivs = [(p, q) for p, q in cov if q > a and p < b]
                cur, gaps = a, []
                for p, q in ivs:
                    if p > cur:
                        gaps.append((cur, p))
                    cur = max(cur, q)
                if cur < b:
                    gaps.append((cur, b))
                for g0, g1 in gaps:
                    if g1 - g0 < 50:
                        continue
                    f0 = g0 - l0 if g0 > a + 1 else g0          # run into the band -> lap
                    f1 = g1 + l0 if g1 < b - 1 else g1
                    f0, f1 = max(f0, a), min(f1, b)
                    for pa, pb in split_bar(f0, f1, zones, stock, l0, 0.0, f"DIST T{dirn}"):
                        out.append(Bar("DIST", "T", dirn, pos, pa, pb, dia, s))
    return out


def link_top_chains(bars):
    """Band bars and distribution bars lying on the same line and overlapping form one lapped run,
    so the plan shows them together with a dimension on every lap."""
    lines = defaultdict(list)
    for b in bars:
        if b.face == "T" and b.fam in ("BAND", "DIST"):
            lines[(b.dirn, round(b.across))].append(b)
    for bl in lines.values():
        if not any(b.fam == "DIST" for b in bl):
            continue
        bl.sort(key=lambda x: x.a)
        run = [bl[0]]
        for p, c in zip(bl, bl[1:]):
            if c.a < p.b - 1:
                run.append(c)
            else:
                _set_chain(run)
                run = [c]
        _set_chain(run)


def _set_chain(run):
    if len(run) < 2:
        return
    _CHAIN[0] += 1
    for k, b in enumerate(run):
        b.chain, b.ci = _CHAIN[0], k


def gen_column_bars(G, P, cfg):
    spec = cfg.get("top_column_bars") or {}
    if not spec.get("enabled", True):
        return []
    bars = []
    f = float(spec.get("extension_factor", 0.25))
    emin = float(spec.get("min_extension", 1000))
    for col in G.columns:
        cs = spec.get(col.kind)
        if not cs:
            continue
        dia, s = int(cs["dia"]), float(cs["spacing"])
        for dirn in ("x", "y"):
            other = "y" if dirn == "x" else "x"
            c_al = col.cx if dirn == "x" else col.cy
            c_ac = col.cy if dirn == "x" else col.cx
            prev, nxt = neighbours(G.sup_lines[dirn], c_al)
            el, er = ext_len(prev, c_al, f, emin), ext_len(nxt, c_al, f, emin)
            n = cs.get("count", "auto")
            if n in (None, "auto"):
                p2, n2 = neighbours(G.sup_lines[other], c_ac)
                spans = [abs(c_ac - v) for v in (p2, n2) if v is not None] or [5000.0]
                width = 0.25 * sum(spans) / len(spans) * (2 if col.kind == "internal" else 1)
                n = max(2, int(round(width / s)) + 1)
            n = int(n)
            base = top_base(cfg, dirn)
            lo, _ = across_range(G, dirn)
            if base:
                off = lo + base / 2
                k0 = round((c_ac - off) / s - (n - 1) / 2)
                poss = [off + (k0 + j) * s for j in range(n)]
            else:
                poss = [c_ac + (j - (n - 1) / 2) * s for j in range(n)]
            for pos in poss:
                for (a, b) in clip_segments(G.inner, dirn, pos, c_al - el, c_al + er, minlen=4 * dia):
                    if a <= c_al + 2 * P.d and b >= c_al - 2 * P.d:
                        bars.append(Bar("COL", "T", dirn, pos, a, b, dia, s, tag=f"C{col.id}"))
    return bars


def gen_wall_bars(G, P, cfg):
    spec = cfg.get("top_wall_bars") or {}
    if not spec.get("enabled"):
        return []
    bars = []
    dia, s = int(spec["dia"]), float(spec["spacing"])
    f = float(spec.get("extension_factor", 0.25))
    emin = float(spec.get("min_extension", 1000))
    gb = cfg.get("top_grid_bands") or {}
    for w in G.walls:
        if w.orient is None:
            continue
        dirn = "x" if w.orient == "y" else "y"         # bars cross the wall
        c_al = w.cx if dirn == "x" else w.cy
        if spec.get("skip_on_grid_bands", True) and top_layout(cfg) != "mesh" and gb.get(dirn):
            if any(abs(c_al - L) < w.thk / 2 + 150 for L in G.col_lines[dirn]):
                continue
        prev, nxt = neighbours(G.sup_lines[dirn], c_al, tol=w.thk / 2 + 150)
        el = w.thk / 2 + ext_len(prev, c_al, f, emin)
        er = w.thk / 2 + ext_len(nxt, c_al, f, emin)
        lo, hi = w.lo + P.cs, w.hi - P.cs
        if hi - lo < s:
            continue
        n = int(math.floor((hi - lo) / s)) + 1
        start = (lo + hi) / 2 - (n - 1) * s / 2
        for j in range(n):
            pos = start + j * s
            for (a, b) in clip_segments(G.inner, dirn, pos, c_al - el, c_al + er, minlen=4 * dia):
                if a <= c_al + w.thk / 2 and b >= c_al - w.thk / 2:
                    bars.append(Bar("WALL", "T", dirn, pos, a, b, dia, s, tag=f"W{w.id}"))
    return bars


def gen_integrity(G, P, cfg):
    spec = cfg.get("integrity_bars") or {}
    if not spec.get("enabled"):
        return []
    bars = []
    kinds = spec.get("columns", ["internal"])
    dia, n, s = int(spec["dia"]), int(spec["count"]), float(spec["spacing"])
    e = spec.get("extension_from_face", "auto")
    ext = ceil50(2 * P.d + P.lbd(dia, "B")) if e in (None, "auto") else float(e)
    for col in G.columns:
        if col.kind not in kinds:
            continue
        for dirn in ("x", "y"):
            c_al = col.cx if dirn == "x" else col.cy
            c_ac = col.cy if dirn == "x" else col.cx
            w_al = col.wx if dirn == "x" else col.wy
            w_ac = col.wy if dirn == "x" else col.wx
            s_use = s
            if n > 1 and (n - 1) * s > w_ac - dia - 50:      # keep 25 mm clear to the column faces
                s_use = math.floor((w_ac - dia - 50) / (n - 1) / 5) * 5
                if s_use < dia + 20:
                    warn(f"integrity bars: {n} bars do not fit through column {col.id} ({w_ac:.0f} mm) - "
                         f"reduce the count")
                    s_use = s
                else:
                    warn(f"integrity bars at column {col.id} ({w_ac:.0f} mm wide): spacing reduced to "
                         f"{s_use:.0f} mm so that all {n} bars pass through the column")
            for j in range(n):
                pos = c_ac + (j - (n - 1) / 2) * s_use
                for (a, b) in clip_segments(G.inner, dirn, pos, c_al - w_al / 2 - ext, c_al + w_al / 2 + ext):
                    if a <= c_al <= b:
                        bars.append(Bar("INTEG", "B", dirn, pos, a, b, dia, s, tag=f"C{col.id}"))
    return bars


def gen_trimmers(G, P, cfg):
    spec = cfg.get("opening_trimmers") or {}
    if not spec.get("enabled"):
        return []
    bars = []
    dia, n, st = int(spec["dia"]), int(spec["count"]), float(spec["spacing"])
    for face in spec.get("faces", ["B", "T"]):
        lb = P.lbd(dia, face)
        for k, hole in enumerate(G.holes):
            minx, miny, maxx, maxy = hole.bounds
            sides = [("x", miny - P.cs - dia, -1, minx, maxx), ("x", maxy + P.cs + dia, 1, minx, maxx),
                     ("y", minx - P.cs - dia, -1, miny, maxy), ("y", maxx + P.cs + dia, 1, miny, maxy)]
            for dirn, base, sign, lo, hi in sides:
                for j in range(n):
                    pos = base + sign * j * st
                    for (a, b) in clip_segments(G.inner, dirn, pos, lo - lb, hi + lb):
                        if a < hi and b > lo:
                            bars.append(Bar("TRIM", face, dirn, pos, a, b, dia, st, tag=f"O{k + 1}"))
    return bars


def add_hooks(bars, G, P):
    """Top bars ending at a slab edge or opening get a bend down (vertical leg)."""
    if not P.hooks:
        return
    edge = G.slab.boundary
    for b in bars:
        if b.face != "T":
            continue
        for end in ("a", "b"):
            p = Point(*pt(b.dirn, getattr(b, end), b.across))
            if p.distance(edge) <= P.cs + 5:
                setattr(b, "hook_" + end, P.leg)


# =============================================================================
# Groups (what is drawn and scheduled)
# =============================================================================
@dataclass
class Group:
    kind: str           # 'bar', 'ubar', 'perim'
    fam: str
    face: str
    dirn: str
    dia: int
    n: int
    step: float
    length: float       # developed length of one bar
    bars: list = field(default_factory=list)
    layer: str = ""
    mark: str = ""
    shape: str = "straight"
    geom: dict = field(default_factory=dict)


FAM_COLORS = {"MESH": 5, "BAND": 1, "DIST": 5, "COL": 6, "WALL": 30, "INTEG": 200, "TRIM": 130, "UBAR": 3,
              "PUNCH": 8}


def layer_for(fam, face, dirn, prefix):
    fn = "BOT" if face == "B" else "TOP"
    if fam in ("UBAR", "PUNCH"):
        return f"{prefix}{fam}"
    if fam == "TRIM":
        return f"{prefix}{fn}_TRIM"
    return f"{prefix}{fn}_{fam}_{dirn.upper()}"


GROUP_SHIFT = [300.0]     # max. shift (mm) along the bar between neighbouring bars of one group


def group_bars(bars, prefix):
    buckets = defaultdict(list)
    for b in bars:
        # same bar (family, face, direction, Ø, length, bends): bars that are shifted a little along
        # their axis (e.g. following a sloping edge) still belong to one group
        key = (b.fam, b.face, b.dirn, b.dia, round(b.length / 10), round(b.hook_a), round(b.hook_b), b.tag)
        buckets[key].append(b)
    groups = []
    for key, bl in buckets.items():
        bl.sort(key=lambda x: (x.across, x.a))
        dedup = [bl[0]]
        for b in bl[1:]:
            if b.across - dedup[-1].across > 1 or abs(b.a - dedup[-1].a) > 1:
                dedup.append(b)
        bl = dedup
        acr = sorted({round(b.across) for b in bl})
        gaps = [q - p for p, q in zip(acr, acr[1:])]
        step = min(gaps) if gaps else bl[0].s
        runs = []                                   # several runs can be open at once (bars in a row)
        for cur in bl:
            best = None
            for r in runs:
                last = r[-1]
                if 1 < cur.across - last.across <= step * 1.05 + 1 and abs(cur.a - last.a) <= GROUP_SHIFT[0]:
                    if best is None or abs(cur.a - last.a) < abs(cur.a - best[-1].a):
                        best = r
            if best is not None:
                best.append(cur)
            else:
                runs.append([cur])
        for r in runs:
            groups.append(make_group(r, step if len(r) > 1 else r[0].s, prefix))
    return groups


MERGE_MAX = [6]


def merge_variable(groups, prefix):
    """Bars that share one end but whose other end changes from bar to bar (a sloping edge,
    a rotated column...) become ONE group with variable length instead of many single bars."""
    # small groups (a few bars each) next to a sloping edge are merged into one variable-length group
    nmax = MERGE_MAX[0]
    frag = [g for g in groups if g.kind == "bar" and g.n <= nmax]
    keep = [g for g in groups if not (g.kind == "bar" and g.n <= nmax)]
    bars = [b for g in frag for b in g.bars]
    buckets = defaultdict(list)
    for b in bars:
        for end in ("a", "b"):
            k = (b.fam, b.face, b.dirn, b.dia, b.tag, b.hook_a > 0, b.hook_b > 0, end, round(getattr(b, end) / 5))
            buckets[k].append(b)

    def runs_of(bl):
        bl = sorted(bl, key=lambda x: x.across)
        out, cur = [], [bl[0]]
        for p, c in zip(bl, bl[1:]):
            if 1 < c.across - p.across <= c.s * 1.05 + 1:
                cur.append(c)
            elif c.across - p.across > 1:
                out.append(cur)
                cur = [c]
        out.append(cur)
        return out

    cand = []
    for bl in buckets.values():
        if len(bl) >= 3:
            cand += [r for r in runs_of(bl) if len(r) >= 3]
    cand.sort(key=len, reverse=True)
    used = set()
    for r in cand:
        r = [b for b in r if id(b) not in used]
        if len(r) < 3:
            continue
        for rr in runs_of(r):
            if len(rr) < 3:
                continue
            gaps = [q.across - p.across for p, q in zip(rr, rr[1:])]
            g = make_group(rr, min(gaps), prefix)
            lens = [b.length for b in rr]
            if max(lens) - min(lens) > 10:
                g.geom["var"] = (min(lens), max(lens))
                g.length = sum(lens) / len(lens)
            keep.append(g)
            used.update(id(b) for b in rr)
    left = [b for b in bars if id(b) not in used]
    if left:
        keep += group_bars(left, prefix)
    return keep


def make_group(run, step, prefix):
    b0 = run[0]
    hooks = (b0.hook_a > 0) + (b0.hook_b > 0)
    shape = {0: "straight", 1: "L", 2: "U"}[hooks]
    return Group("bar", b0.fam, b0.face, b0.dirn, b0.dia, len(run), step, b0.length, run,
                 layer_for(b0.fam, b0.face, b0.dirn, prefix), shape=shape)


def build_chains(bars, groups):
    """Runs of lapped bars. Runs made of the same bar groups are drawn once, together,
    so the overlaps can be shown and dimensioned."""
    gid = {}
    for g in groups:
        for b in g.bars:
            gid[id(b)] = g
    runs = defaultdict(list)
    for b in bars:
        if b.chain >= 0:
            runs[b.chain].append(b)
    sigs = defaultdict(list)
    for bl in runs.values():
        bl.sort(key=lambda x: x.ci)
        sig = tuple(id(gid[id(b)]) for b in bl)
        sigs[sig].append(bl)
    out = []
    for sig, runs_ in sigs.items():
        runs_.sort(key=lambda bl: bl[0].across)
        out.append({"runs": runs_, "groups": [gid[id(b)] for b in runs_[0]],
                    "face": runs_[0][0].face, "dirn": runs_[0][0].dirn})
    out.sort(key=lambda c: -len(c["runs"]))
    return out


def gen_ubars(G, P, cfg, prefix):
    spec = cfg.get("free_edge_ubars") or {}
    if not spec.get("enabled"):
        return []
    dia, s = int(spec["dia"]), float(spec["spacing"])
    leg = spec.get("leg", "auto")
    leg = 2 * P.h if leg in (None, "auto") else float(leg)
    dev = 2 * leg + (P.h - P.ct - P.cb)
    walls = G.walls_union.buffer(P.cs + 50) if not G.walls_union.is_empty else Polygon()
    groups = []
    rings = []
    for g in getattr(G.slab, "geoms", [G.slab]):
        rings.append(g.exterior)
        rings += list(g.interiors)
    for ring in rings:
        c = list(ring.coords)
        for p, q in zip(c, c[1:]):
            seg = LineString([p, q])
            if seg.length < s:
                continue
            free = seg.difference(walls) if not walls.is_empty else seg
            for piece in getattr(free, "geoms", [free]):
                if piece.is_empty or piece.geom_type != "LineString" or piece.length < s + 2 * P.cs:
                    continue
                (x0, y0), (x1, y1) = piece.coords[0], piece.coords[-1]
                Lp = piece.length
                tx, ty = (x1 - x0) / Lp, (y1 - y0) / Lp
                nx, ny = -ty, tx
                mid = ((x0 + x1) / 2, (y0 + y1) / 2)
                if not G.slab.contains(Point(mid[0] + nx * 20, mid[1] + ny * 20)):
                    nx, ny = -nx, -ny
                n = int(math.floor((Lp - 2 * P.cs) / s)) + 1
                used = (n - 1) * s
                st = (Lp - used) / 2
                pos = [(x0 + tx * (st + i * s), y0 + ty * (st + i * s)) for i in range(n)]
                dirn = "x" if abs(nx) > abs(ny) else "y"
                g_ = Group("ubar", "UBAR", "B", dirn, dia, n, s, dev, layer=f"{prefix}UBAR", shape="U")
                g_.geom = {"pos": pos, "t": (tx, ty), "nrm": (nx, ny), "leg": leg, "cs": P.cs}
                groups.append(g_)
    return groups


def gen_perimeters(G, P, cfg, prefix):
    spec = cfg.get("punching_perimeters") or {}
    if not spec.get("enabled"):
        return []
    kinds = spec.get("columns", ["internal", "edge", "corner"])
    out = []
    for col in G.columns:
        if col.kind not in kinds:
            continue
        u1 = col.poly.buffer(2 * P.d, resolution=16)
        line = u1.exterior.intersection(G.slab)
        g = Group("perim", "PUNCH", "T", "x", 0, 0, 0, 0, layer=f"{prefix}PUNCH")
        g.geom = {"line": line, "col": col}
        out.append(g)
    return out


def assign_marks(groups):
    def key(g):
        v = g.geom.get("var")
        return (g.dia, int(round(g.length / 10) * 10), g.shape, (int(v[0]), int(v[1])) if v else (0, 0))
    keys = sorted({key(g) for g in groups if g.kind in ("bar", "ubar")})
    marks = {k: f"N{i + 1}" for i, k in enumerate(keys)}
    for g in groups:
        if g.kind in ("bar", "ubar"):
            k = key(g)
            g.length = round(g.length / 10) * 10
            g.mark = marks[k]
    return marks


# =============================================================================
# Drawing
# =============================================================================
class Drawer:
    def __init__(self, doc, cfg, P, G, off=(0.0, 0.0)):
        self.doc, self.cfg, self.P, self.G = doc, cfg, P, G
        self.off = off                    # where this plan sits in the drawing (mm)
        dr = cfg.get("drawing", {})
        self.bar_color = int(dr.get("bar_color", 7))
        self.bar_width = float(dr.get("bar_width", 20))
        self.text_color = int(dr.get("text_color", 7))
        self.show_dims = bool(dr.get("lap_dimensions", False))
        self.show_axis_dims = bool(dr.get("axis_to_bar_end_dimensions", True))
        self.show_dist = bool(dr.get("distribution_lines", False))
        self.show_notes = bool(dr.get("notes", False))
        self.bar_style = dr.get("bar_style", "dimension")          # dimension | polyline
        self.dim_style = str(dr.get("dim_style_name", "REBAR M-200"))
        self.dim_color = int(dr.get("dim_color", 6))
        self.dim_lw = int(dr.get("dim_lineweight", 53))
        self.dim_text_color = int(dr.get("dim_text_color", 7))
        self.labelled = set()             # groups whose label is already on the drawing
        # label_mode: length = every bar dimension shows only its own measured length
        #             full   = mark, number of bars, Ø, length, spacing  (e.g. N23 39ф12x<>/20)
        self.label_mode = dr.get("label_mode", "length")
        # text under every drawn bar: number_of_bars Ø dia x length / spacing  (e.g. 7Φ12x458/20)
        self.show_notes_bar = bool(dr.get("bar_notes", True))
        self.noted = set()
        self.u = G.u
        self.th = float(dr.get("text_height", 150))
        self.mode = dr.get("mode", "representative")
        self.sym = dr.get("dia_symbol", "ф")
        self.fmt = dr.get("label_format", "{mark} {n}{sym}{d}x{L}/{s}")
        self.msp = doc.modelspace()
        self.placed = []
        self.forced = {}                  # group id -> across of its drawn (lapped) bar
        self.chain_drawn = set()          # groups whose representative bar is drawn by a chain
        if "REBAR_TXT" not in doc.styles:
            doc.styles.add("REBAR_TXT", font="arial.ttf")
        dim_unit = dr.get("dimension_unit", "cm")
        self.dimlfac = self.u / UNIT_MM.get(dim_unit, 10.0)
        if "REBAR_DIM" not in doc.dimstyles:
            t = self.th / self.u
            doc.dimstyles.new("REBAR_DIM", dxfattribs={
                "dimtxt": 0.8 * t, "dimasz": 0.5 * t, "dimtsz": 0.35 * t, "dimexo": 0.15 * t,
                "dimexe": 0.3 * t, "dimgap": 0.15 * t, "dimlfac": self.dimlfac, "dimdec": 0,
                "dimtad": 1, "dimtih": 0, "dimtoh": 0, "dimzin": 8, "dimtxsty": "REBAR_TXT"})
        if self.bar_style == "dimension" and self.dim_style not in doc.dimstyles:
            # bar = one dimension: magenta dimension line, no arrowheads, both extension lines
            # suppressed, measured length as live text (proportions as the M-200 style)
            t = self.th / self.u
            ds = doc.dimstyles.new(self.dim_style, dxfattribs={
                "dimclrd": self.dim_color, "dimlwd": self.dim_lw, "dimdle": 0, "dimdli": 5 * t,
                "dimclre": self.dim_color, "dimse1": 1, "dimse2": 1, "dimexo": t * 12.5 / 15, "dimexe": t * 25 / 15,
                "dimlwe": -2, "dimasz": t, "dimtsz": 0, "dimcen": t * 50 / 15,
                "dimtxsty": "Standard" if "Standard" in doc.styles else "REBAR_TXT",
                "dimclrt": self.dim_text_color, "dimtxt": t, "dimtad": 1, "dimgap": t * 10 / 15, "dimjust": 0,
                "dimtih": 0, "dimtoh": 0, "dimtmove": 2, "dimtofl": 1, "dimatfit": 3,
                "dimlunit": 2, "dimdec": 0, "dimzin": 8, "dimlfac": self.dimlfac, "dimscale": 1})
            ds.set_arrows(blk=ARROWS.none, ldrblk=ARROWS.none)
        if self.show_dims and "REBAR LAP" not in doc.dimstyles:
            # lap length: short dimension over the two overlapping bars, oblique ticks, live length
            t = self.th / self.u
            doc.dimstyles.new("REBAR LAP", dxfattribs={
                "dimclrd": self.dim_color, "dimlwd": 25, "dimclre": self.dim_color, "dimlwe": 18,
                "dimse1": 0, "dimse2": 0, "dimexo": 0, "dimexe": 0.3 * t, "dimdle": 0,
                "dimtsz": 0.4 * t, "dimtxsty": "Standard" if "Standard" in doc.styles else "REBAR_TXT",
                "dimclrt": self.dim_text_color, "dimtxt": 0.9 * t, "dimtad": 1, "dimgap": 0.3 * t,
                "dimjust": 0, "dimtih": 0, "dimtoh": 0, "dimtmove": 2, "dimtofl": 1, "dimatfit": 3,
                "dimlunit": 2, "dimdec": 0, "dimzin": 8, "dimlfac": self.dimlfac, "dimscale": 1})
        self.show_range = bool(dr.get("distribution_dimensions", True))
        if self.show_range and "REBAR RANGE" not in doc.dimstyles:
            # distribution range of a bar group: from the first to the last bar, oblique ticks
            t = self.th / self.u
            doc.dimstyles.new("REBAR RANGE", dxfattribs={
                "dimclrd": self.dim_color, "dimlwd": 25, "dimclre": self.dim_color, "dimse1": 1, "dimse2": 1,
                "dimtsz": 0.5 * t, "dimdle": 0, "dimexo": 0, "dimexe": 0,
                "dimtxsty": "Standard" if "Standard" in doc.styles else "REBAR_TXT",
                "dimclrt": self.dim_text_color, "dimtxt": t, "dimtad": 1, "dimgap": t * 10 / 15, "dimjust": 0,
                "dimtih": 0, "dimtoh": 0, "dimtmove": 2, "dimtofl": 1, "dimatfit": 3,
                "dimlunit": 2, "dimdec": 0, "dimzin": 8, "dimlfac": self.dimlfac, "dimscale": 1})
    def m(self, p):                       # mm -> drawing units (+ plan offset)
        return ((p[0] + self.off[0]) / self.u, (p[1] + self.off[1]) / self.u)

    def layer(self, name, fam):
        if name not in self.doc.layers:
            self.doc.layers.add(name, color=FAM_COLORS.get(fam, 7))

    def line(self, pts, layer, lw=25):
        self.msp.add_lwpolyline([self.m(p) for p in pts], dxfattribs={"layer": layer, "lineweight": lw})

    def rebar(self, pts, layer, text=" ", text_at=None, note=None):
        """One reinforcement bar.
        bar_style = dimension: a single DIMENSION whose dimension line is the bar (magenta, no arrows,
        extension lines suppressed); its text is the bar label with the live length (<>), so the bar
        can be stretched in AutoCAD and the length updates by itself.
        bar_style = polyline: white polyline with a constant width."""
        if self.bar_style != "dimension":
            self.msp.add_lwpolyline([self.m(p) for p in pts],
                                    dxfattribs={"layer": layer, "color": self.bar_color,
                                                "const_width": self.bar_width / self.u})
            if text.strip():
                self._poly_label(pts, layer, text, text_at)
            return
        # straight part (longest segment) = the dimension; bent legs (if any) as short magenta lines
        segs = list(zip(pts, pts[1:]))
        i0 = max(range(len(segs)), key=lambda i: math.dist(*segs[i]))
        p1, p2 = segs[i0]
        for i, (q1, q2) in enumerate(segs):
            if i != i0:
                self.msp.add_line(self.m(q1), self.m(q2),
                                  dxfattribs={"layer": layer, "color": self.dim_color, "lineweight": self.dim_lw})
        if self.label_mode == "length":
            text, text_at = "<>", None                 # just the live length of the dimension line
        ang = math.degrees(math.atan2(p2[1] - p1[1], p2[0] - p1[0]))
        try:
            dim = self.msp.add_linear_dim(base=self.m(p1), p1=self.m(p1), p2=self.m(p2), angle=ang,
                                          dimstyle=self.dim_style, text=text, dxfattribs={"layer": layer})
            if text_at is not None:
                dim.set_location(self.m(text_at), leader=False, relative=False)
            dim.render()
            if dim.dimension.dxf.get("text_midpoint") is None:      # suppressed text: keep a valid entity
                mp = self.m(((p1[0] + p2[0]) / 2, (p1[1] + p2[1]) / 2))
                dim.dimension.dxf.text_midpoint = (mp[0], mp[1], 0)
        except Exception as ex:
            warn(f"bar dimension could not be drawn: {ex}")
        if note:
            self.note_under(p1, p2, note, layer)

    def note_under(self, p1, p2, note, layer):
        """Bar description written under the bar (right of a vertical bar), centred on it."""
        gap = self.th * 10 / 15
        horiz = abs(p2[0] - p1[0]) >= abs(p2[1] - p1[1])
        mx, my = (p1[0] + p2[0]) / 2, (p1[1] + p2[1]) / 2
        try:
            if horiz:
                t = self.msp.add_text(note, height=self.th / self.u, dxfattribs={
                    "layer": layer + "_TEXT", "style": "REBAR_TXT", "color": self.dim_text_color})
                t.set_placement(self.m((mx, my - gap)), align=ezdxf.enums.TextEntityAlignment.TOP_CENTER)
            else:
                t = self.msp.add_text(note, height=self.th / self.u, rotation=90, dxfattribs={
                    "layer": layer + "_TEXT", "style": "REBAR_TXT", "color": self.dim_text_color})
                t.set_placement(self.m((mx + gap, my)), align=ezdxf.enums.TextEntityAlignment.TOP_CENTER)
        except Exception as ex:
            warn(f"bar note could not be written: {ex}")

    def _poly_label(self, pts, layer, text, text_at):
        (x0, y0), (x1, y1) = pts[0], pts[-1]
        horiz = abs(x1 - x0) >= abs(y1 - y0)
        txt = text.replace("<>", f"{math.hypot(x1 - x0, y1 - y0) / 10:.0f}")
        at = text_at or ((x0 + x1) / 2, (y0 + y1) / 2)
        if horiz:
            self.text(txt, (at[0], at[1] + 0.25 * self.th), 0, layer)
        else:
            self.text(txt, (at[0] - 0.25 * self.th, at[1]), 90, layer)

    def text(self, txt, at, rot, layer, h=None):
        h = h or self.th
        t = self.msp.add_text(txt, height=h / self.u, rotation=rot,
                              dxfattribs={"layer": layer, "style": "REBAR_TXT", "color": self.text_color})
        t.set_placement(self.m(at), align=ezdxf.enums.TextEntityAlignment.BOTTOM_CENTER)

    def label(self, g, live=False):
        """Bar label. live=True puts <> for the length, so a bar-dimension shows its own measured length."""
        v = g.geom.get("var")
        has_legs = g.shape != "straight"
        if v:
            L = f"({v[0] / 10:.0f}÷{v[1] / 10:.0f})"
        elif live and not has_legs:
            L = "<>"
        else:
            L = f"{g.length / 10:.0f}"
        return self.fmt.format(mark=g.mark, n=g.n, sym=self.sym, d=g.dia, L=L, s=f"{g.step / 10:.0f}")

    def free(self, poly):
        return not any(poly.intersects(q) for q in self.placed)

    def bar_pts(self, br, ac):
        """Plan polyline of one bar at across-position ac, bends drawn as stubs."""
        d = br.dirn
        side = -1 if d == "x" else 1
        pts = []
        if br.hook_a:
            pts.append(pt(d, br.a, ac + side * br.hook_a))
        pts += [pt(d, br.a, ac), pt(d, br.b, ac)]
        if br.hook_b:
            pts.append(pt(d, br.b, ac + side * br.hook_b))
        return pts

    def draw_chain(self, ch):
        """Draw one run of lapped bars: consecutive bars offset a little so the overlap is
        visible, and every overlap dimensioned (lap length)."""
        runs = ch["runs"]
        d = ch["dirn"]
        off = 0.45 * self.th                      # visual offset between the two lapped bars
        lab = 1 if d == "x" else -1              # labels above X-bars, left of Y-bars
        a0, b1 = runs[0][0].a, runs[0][-1].b
        # choose a run (across position) whose label/dimension band is still free
        # prefer runs in the middle of a panel, away from columns and labels over supports
        other = "y" if d == "x" else "x"
        sl = self.G.sup_lines[other]
        mids = [(p + q) / 2 for p, q in zip(sl, sl[1:])] or [runs[len(runs) // 2][0].across]
        order = sorted(runs, key=lambda bl: min(abs(bl[0].across - m_) for m_ in mids))
        chosen = None
        for bl in order:
            ac = bl[0].across
            lo, hi = (ac - 2.2 * self.th, ac + off + 0.2 * self.th) if lab == 1 else \
                     (ac - 0.2 * self.th, ac + off + 2.2 * self.th)
            bx = box(a0, lo, b1, hi) if d == "x" else box(lo, a0, hi, b1)
            if self.free(bx):
                chosen = (bl, bx)
                break
        if chosen is None:
            chosen = (order[0], None)
        bl, bx = chosen
        if bx is not None:
            self.placed.append(bx)
        ac0 = bl[0].across
        for k, br in enumerate(bl):
            g = ch["groups"][k]
            ac = ac0 + (off if k % 2 else 0.0)
            self.layer(g.layer, g.fam)
            txt = " "
            if id(g) not in self.labelled:
                txt = self.label(g, live=True)
                self.labelled.add(id(g))
            self.layer(g.layer + "_TEXT", g.fam)
            note = None
            if self.show_notes_bar and id(g) not in self.noted:      # one note per bar group
                note = self.bar_note(g)
                self.noted.add(id(g))
            self.rebar(self.bar_pts(br, ac), g.layer, text=txt, note=note)
            self.chain_drawn.add(id(g))
            self.forced.setdefault(id(g), ac)
        # lap dimensions
        for k in range(len(bl) - 1 if self.show_dims else 0):
            cur, nxt = bl[k], bl[k + 1]
            if cur.b <= nxt.a:
                continue
            ac_c = ac0 + (off if k % 2 else 0.0)
            ac_n = ac0 + (off if (k + 1) % 2 else 0.0)
            # dimension line just over the two overlapping bars (above X bars, left of Y bars),
            # its ends exactly at the start of the next bar and the end of the current bar
            base = ac0 + off + 0.8 * self.th if d == "x" else ac0 - 0.8 * self.th
            p1 = self.m(pt(d, nxt.a, ac_n))
            p2 = self.m(pt(d, cur.b, ac_c))
            layer = ch["groups"][k].layer + "_LAP"
            self.layer(layer, ch["groups"][k].fam)
            try:
                dim = self.msp.add_linear_dim(base=self.m(pt(d, (nxt.a + cur.b) / 2, base)), p1=p1, p2=p2,
                                              angle=0 if d == "x" else 90, dimstyle="REBAR LAP",
                                              dxfattribs={"layer": layer})
                dim.render()
                if self.show_axis_dims:
                    self.axis_to_end(d, cur.b, ac_c, ac0, off, ch["groups"][k])
            except Exception as ex:
                warn(f"lap dimension could not be drawn: {ex}")

    def axis_coords(self, d):
        """Positions of the axes crossing bars of direction d (vertical axes for X bars)."""
        ax = self.G.axes["x" if d == "x" else "y"]
        if ax and self.G.axes["y" if d == "x" else "x"]:
            return [a[0] for a in ax]
        return list(self.G.sup_lines["x" if d == "x" else "y"])   # same fallback as the formwork plan

    def axis_to_end(self, d, end, ac_bar, ac0, off, g):
        """Dimension from the closest axis to the end of a bar (where it laps), drawn on the other side
        of the bars than the lap dimension: below X bars, right of Y bars."""
        axes = self.axis_coords(d)
        if not axes:
            return
        ax = min(axes, key=lambda v: abs(v - end))
        if abs(ax - end) < 10:                       # bar ends on the axis
            return
        base = ac0 - 1.7 * self.th if d == "x" else ac0 + off + 1.7 * self.th   # room for the text
        layer = g.layer + "_AXDIM"
        self.layer(layer, g.fam)
        try:
            dim = self.msp.add_linear_dim(base=self.m(pt(d, (ax + end) / 2, base)),
                                          p1=self.m(pt(d, ax, ac_bar)), p2=self.m(pt(d, end, ac_bar)),
                                          angle=0 if d == "x" else 90, dimstyle="REBAR LAP",
                                          dxfattribs={"layer": layer})
            dim.render()
        except Exception as ex:
            warn(f"axis dimension could not be drawn: {ex}")

    def plan_ranges(self, groups):
        """Distribution dimensions of one plan. Each group covers from half a spacing before its first
        bar to half a spacing after its last bar (cut at the slab edge / opening), so the dimensions of
        neighbouring groups meet exactly. Neighbouring groups get the same position along the bars,
        so their dimensions form one continuous chain."""
        self.ranges = {}
        gs = [g for g in groups if g.kind == "bar" and len(g.bars) >= 2]
        info = {}
        for g in gs:
            bars = g.bars
            lo_c, hi_c = max(b.a for b in bars), min(b.b for b in bars)
            if hi_c <= lo_c:
                br = bars[len(bars) // 2]
                lo_c, hi_c = br.a, br.b
            info[id(g)] = (g, bars[0].across, bars[-1].across, lo_c, hi_c)
        # neighbours: same direction, next group starts one spacing after this one ends, along ranges overlap
        adj = defaultdict(set)
        items = sorted(info.values(), key=lambda t: t[1])
        for i, (g, p0, p1, l0, h0) in enumerate(items):
            for (g2, q0, q1, l1, h1) in items[i + 1:]:
                if q0 > p1 + 1.5 * g.step:
                    break
                if g2.dirn == g.dirn and abs(q0 - p1 - g.step) < 0.25 * g.step and min(h0, h1) > max(l0, l1):
                    adj[id(g)].add(id(g2))
                    adj[id(g2)].add(id(g))
        seen = set()
        for k in info:
            if k in seen:
                continue
            comp, stack = [], [k]
            while stack:
                c = stack.pop()
                if c in seen:
                    continue
                seen.add(c)
                comp.append(c)
                stack += list(adj[c] - seen)
            # walk the chain across the bars; start a new straight run where no common place is left
            comp.sort(key=lambda c: (info[c][1], info[c][3], info[c][4], info[c][0].mark))   # no ties by address
            run, lo, hi = [], None, None
            for c in comp + [None]:
                if c is not None:
                    nlo, nhi = (info[c][3], info[c][4]) if lo is None else (max(lo, info[c][3]), min(hi, info[c][4]))
                    if not run or nhi - nlo > 0.2 * (info[c][4] - info[c][3]):
                        run.append(c)
                        lo, hi = nlo, nhi
                        continue
                along_run = lo + 0.25 * (hi - lo)
                for r in run:
                    self.ranges[r] = along_run
                if c is not None:
                    run, lo, hi = [c], info[c][3], info[c][4]

    def range_extent(self, g: Group):
        """(along, lo, hi) of the distribution dimension of a group: from half a spacing before the
        first bar to half a spacing after the last one, cut at the slab / opening edge."""
        bars = g.bars
        along = getattr(self, "ranges", {}).get(id(g))
        if along is None:
            lo_c, hi_c = max(b.a for b in bars), min(b.b for b in bars)
            along = lo_c + 0.25 * (hi_c - lo_c) if hi_c > lo_c else (bars[0].a + bars[0].b) / 2
        p0, p1 = bars[0].across, bars[-1].across
        lo, hi = p0 - g.step / 2, p1 + g.step / 2
        G = self.G
        x0, y0, x1, y1 = G.slab.bounds
        ln = LineString([pt(g.dirn, along, y0 - 10), pt(g.dirn, along, y1 + 10)]) if g.dirn == "x" else \
            LineString([pt(g.dirn, along, x0 - 10), pt(g.dirn, along, x1 + 10)])
        inter = ln.intersection(G.slab)
        mid = (p0 + p1) / 2
        for part in getattr(inter, "geoms", [inter]):
            if part.geom_type != "LineString":
                continue
            cs = [c[1] if g.dirn == "x" else c[0] for c in part.coords]
            if min(cs) - 1 <= mid <= max(cs) + 1:
                lo, hi = max(lo, min(cs)), min(hi, max(cs))
                break
        return along, lo, hi

    def bar_note(self, g: Group):
        """Text under a bar:  number_of_bars Ø dia x length / spacing   (length and spacing in cm).
        Number of bars = width of the reinforced area (distribution dimension) / spacing."""
        if len(g.bars) >= 2:
            _, lo, hi = self.range_extent(g)
            n = max(1, math.ceil((hi - lo) / g.step - 0.01))
        else:
            n = 1
        v = g.geom.get("var")
        L = f"({v[0] / 10:.0f}÷{v[1] / 10:.0f})" if v else f"{g.length / 10:.0f}"
        return f"{n}{self.sym}{g.dia}x{L}/{g.step / 10:.0f}"

    def draw_range(self, g: Group):
        """Dimension across a bar group = width of the area reinforced by that bar."""
        bars = g.bars
        if not self.show_range or len(bars) < 2:
            return
        along, lo, hi = self.range_extent(g)
        try:
            dim = self.msp.add_linear_dim(base=self.m(pt(g.dirn, along, lo)), p1=self.m(pt(g.dirn, along, lo)),
                                          p2=self.m(pt(g.dirn, along, hi)), angle=90 if g.dirn == "x" else 0,
                                          dimstyle="REBAR RANGE", dxfattribs={"layer": g.layer + "_RANGE"})
            dim.render()
        except Exception as ex:
            warn(f"distribution dimension could not be drawn: {ex}")

    def draw_bar_group(self, g: Group):
        self.layer(g.layer, g.fam)
        if self.show_range:
            self.layer(g.layer + "_RANGE", g.fam)
            self.draw_range(g)
        if id(g) in self.chain_drawn and self.bar_style == "dimension":
            return                                     # drawn and labelled with its lapped run
        txt = self.label(g, live=True)
        tw = 0.62 * self.th * len(txt)
        bars = g.bars
        n = len(bars)
        a, b = bars[0].a, bars[0].b
        forced = self.forced.get(id(g))
        lo_c = max(br.a for br in bars)                # part common to all bars of the group
        hi_c = min(br.b for br in bars)
        if hi_c <= lo_c:
            lo_c, hi_c = a, b
        # pick a representative bar and label position that do not collide with other labels
        idx_opts = [n // 2, n // 4, (3 * n) // 4, n // 8, (7 * n) // 8, 0, n - 1]
        t_opts = [0.5, 0.35, 0.65, 0.2, 0.8]
        chosen = None
        for ti in t_opts:
            for ii in idx_opts:
                ac = bars[min(ii, n - 1)].across if forced is None else forced
                mid = a + ti * (b - a)
                if g.dirn == "x":
                    bx = box(mid - tw / 2, ac + 0.2 * self.th, mid + tw / 2, ac + 1.9 * self.th)
                else:
                    bx = box(ac - 1.9 * self.th, mid - tw / 2, ac - 0.2 * self.th, mid + tw / 2)
                if self.free(bx):
                    chosen = (ii, ti, bx)
                    break
            if chosen:
                break
        if not chosen:
            ii, ti = n // 2, 0.5
            ac = bars[ii].across if forced is None else forced
            mid = a + 0.5 * (b - a)
            chosen = (ii, ti, box(mid - 1, ac - 1, mid + 1, ac + 1))
        ii, ti, bx = chosen
        self.placed.append(bx)
        ac = bars[min(ii, n - 1)].across if forced is None else forced
        # representative bar (with legs drawn as stubs) - lapped bars are already drawn by their chain
        rep_bar = bars[min(ii, n - 1)]
        a, b = rep_bar.a, rep_bar.b
        mid = a + ti * (b - a)
        gap = self.th * 10 / 15 + 0.5 * self.th
        text_at = None
        if abs(ti - 0.5) > 1e-6:                     # label moved along the bar to avoid a clash
            text_at = pt(g.dirn, mid, ac + gap if g.dirn == "x" else ac - gap)
        if id(g) not in self.chain_drawn:
            self.layer(g.layer + "_TEXT", g.fam)
            note = None
            if self.show_notes_bar and id(g) not in self.noted:
                note = self.bar_note(g)
                self.noted.add(id(g))
            self.rebar(self.bar_pts(rep_bar, ac), g.layer, text=txt, text_at=text_at, note=note)
        else:                                          # polyline style: label of a lapped group
            self._poly_label([pt(g.dirn, a, ac), pt(g.dirn, b, ac)], g.layer,
                             txt.replace("<>", f"{g.length / 10:.0f}"), pt(g.dirn, mid, ac))
        # other bars of the group (no text): all of them in mode 'all', else shortest + longest of a
        # variable-length group
        extra = bars if self.mode == "all" else ((bars[0], bars[-1]) if g.geom.get("var") else ())
        for br in extra:
            if not (abs(ac - br.across) < 1):
                self.rebar(self.bar_pts(br, br.across), g.layer)
        # distribution line across the group with end ticks
        if n > 1 and self.show_dist:
            along = lo_c + (0.15 if ti >= 0.35 else 0.85) * (hi_c - lo_c)
            p0, p1 = bars[0].across, bars[-1].across
            self.line([pt(g.dirn, along, p0), pt(g.dirn, along, p1)], g.layer, lw=18)
            k = 0.35 * self.th
            for pp in (p0, p1):
                c = pt(g.dirn, along, pp)
                self.line([(c[0] - k, c[1] - k), (c[0] + k, c[1] + k)], g.layer, lw=18)

    def draw_ubar_group(self, g: Group):
        self.layer(g.layer, "UBAR")
        G = g.geom
        (tx, ty), (nx, ny), leg, cs = G["t"], G["nrm"], G["leg"], G["cs"]
        pos = G["pos"]
        if self.mode == "all":
            for p in pos:
                self.rebar([(p[0] + nx * cs, p[1] + ny * cs), (p[0] + nx * (cs + leg), p[1] + ny * (cs + leg))],
                           g.layer)
        p = pos[len(pos) // 2]
        self.rebar([(p[0] + nx * cs, p[1] + ny * cs), (p[0] + nx * (cs + leg), p[1] + ny * (cs + leg))], g.layer)
        q0, q1 = pos[0], pos[-1]
        off = cs + 0.5 * leg
        if self.show_dist:
            self.line([(q0[0] + nx * off, q0[1] + ny * off), (q1[0] + nx * off, q1[1] + ny * off)], g.layer, lw=18)
        txt = "U " + self.label(g)
        rot = math.degrees(math.atan2(ty, tx))
        if rot > 90.01 or rot <= -90:
            rot -= 180 if rot > 0 else -180
        ref = pos[max(0, len(pos) // 4)]
        off2 = cs + leg + 0.3 * self.th
        self.text(txt, (ref[0] + nx * off2, ref[1] + ny * off2), rot, g.layer, h=self.th * 0.85)

    def draw_perim(self, g: Group):
        self.layer(g.layer, "PUNCH")
        line = g.geom["line"]
        for part in getattr(line, "geoms", [line]):
            if part.geom_type == "LineString" and part.length > 10:
                self.msp.add_lwpolyline([self.m(c) for c in part.coords],
                                        dxfattribs={"layer": g.layer, "linetype": "DASHED", "lineweight": 18})
        col = g.geom["col"]
        r = max(col.wx, col.wy) / 2 + 2 * self.P.d
        self.text(f"u1 (2d={2 * self.P.d:.0f})", (col.cx + 0.6 * r, col.cy + 0.75 * r), 0, g.layer, h=self.th * 0.7)

    def notes(self, P, laps, face=None, title=None):
        x0, y0, x1, y1 = self.G.bounds
        if title:
            self.layer("REB_NOTES", "NOTES")
            t = self.msp.add_text(title, height=2.5 * self.th / self.u,
                                  dxfattribs={"layer": "REB_NOTES", "style": "REBAR_TXT", "color": self.text_color})
            ty = y0 - 12 * self.th
            if self.cfg.get("drawing", {}).get("axes_on_plans", True):   # keep clear of the axis bubbles
                r = 2.2 * self.th
                lows = [lo for _, lo, _ in self.G.axes["x"]] or [y0]
                ty = min(y0 - (8 + 3) * self.th - 2 * r, min(lows) - 2 * r) - 3 * self.th
            t.set_placement(self.m(((x0 + x1) / 2, ty)),
                            align=ezdxf.enums.TextEntityAlignment.TOP_CENTER)
        if not self.show_notes:
            return
        lines = ["REINFORCEMENT NOTES (generated)",
                 f"Concrete C{P.fck:.0f}, steel B{P.fyk:.0f}B, h = {P.h:.0f} mm",
                 f"Cover: bottom {P.cb:.0f} mm, top {P.ct:.0f} mm",
                 "Lap lengths l0 (EC2 8.7.3, " + ("staggered" if P.stagger else "all bars lapped at one section") + "):"]
        for d in sorted(laps):
            lb, lt = laps[d]
            if face == "B":
                lines.append(f"  {self.sym}{d}: {lb:.0f} mm")
            elif face == "T":
                lines.append(f"  {self.sym}{d}: {lt:.0f} mm")
            else:
                lines.append(f"  {self.sym}{d}: bottom {lb:.0f} mm, top {lt:.0f} mm")
        if face in (None, "B"):
            lines.append("Bottom bars lapped at mid-span." if is_mat() else "Bottom bars lapped over the supports.")
        if face in (None, "T"):
            lines.append("Top bars lapped over the supports." if is_mat() else "Top bars lapped at mid-span.")
        lines.append("Bar marks: N.. number Ø x length(cm) / spacing(cm). Laps dimensioned in cm.")
        self.layer("REB_NOTES", "NOTES")
        mt = self.msp.add_mtext("\\P".join(lines), dxfattribs={"layer": "REB_NOTES", "style": "REBAR_TXT",
                                                              "char_height": self.th / self.u})
        mt.set_location(self.m((x0, y1 + 6 * self.th + len(lines) * 1.8 * self.th)))



# =============================================================================
# Formwork plan
# =============================================================================
class Formwork:
    """Formwork plan of the slab: hatched walls and columns with their names and sizes, axis bubbles,
    dimension chains, rotated slab sections (R.C. slab section) with the slab thickness, legend, notes."""

    def __init__(self, doc, cfg, P, G, off):
        self.doc, self.cfg, self.P, self.G, self.off = doc, cfg, P, G, off
        dr = cfg.get("drawing", {})
        fw = cfg.get("formwork", {}) or {}
        self.fw = fw
        self.u = G.u
        self.th = float(dr.get("text_height", 150))
        self.msp = doc.modelspace()
        self.tcol = int(fw.get("text_color", 7))
        self.strip_half = defaultdict(float)       # set by slab_sections (formwork plan only)
        self.strip_at = {}
        self.sections = []                         # wall sections: (wall, tc, side list)
        self.ws_cover = Polygon()                  # area covered by the wall sections
        self._cores = None
        # what is already on the plan (texts, sections, walls, columns): later texts keep clear of it
        self.occ = defaultdict(list)
        for g in getattr(G.walls_union, "geoms", [G.walls_union]):
            self._reg(g)
        for c in G.columns:
            self._reg(c.poly)
        self.ws_stub = float((fw.get("wall_sections") or {}).get("wall_extension", 150))
        for name, col in [("FW_WALL_HATCH", 8), ("FW_COLUMN_HATCH", 8), ("FW_SLAB_SECTION", 8),
                          ("FW_WALL_SECTION", 8), ("FW_DIM", 7), ("FW_TEXT", 7), ("FW_AXIS", 7),
                          ("FW_OPENING", 1)]:
            if name not in doc.layers:
                doc.layers.add(name, color=col)
        if "FW_DIM" not in doc.dimstyles:
            t = self.th / self.u
            doc.dimstyles.new("FW_DIM", dxfattribs={
                "dimtxt": 1.2 * t, "dimtsz": 0.6 * t, "dimasz": 0.6 * t, "dimexo": 0.3 * t, "dimexe": 0.5 * t,
                "dimgap": 0.3 * t, "dimdec": 0, "dimzin": 8, "dimtad": 1, "dimtih": 0, "dimtoh": 0,
                "dimlfac": self.u / UNIT_MM.get(dr.get("dimension_unit", "cm"), 10.0), "dimdle": 0.5 * t,
                "dimclrd": 256, "dimclre": 256, "dimclrt": self.tcol, "dimtxsty": "Standard"
                if "Standard" in doc.styles else "REBAR_TXT", "dimatfit": 3, "dimtmove": 2})

    # --- helpers --------------------------------------------------------------
    def m(self, p):
        return ((p[0] + self.off[0]) / self.u, (p[1] + self.off[1]) / self.u)

    def text(self, txt, at, h, rot=0, align="MIDDLE_CENTER", layer="FW_TEXT"):
        t = self.msp.add_text(txt, height=h / self.u, rotation=rot,
                              dxfattribs={"layer": layer, "style": "REBAR_TXT", "color": self.tcol})
        t.set_placement(self.m(at), align=getattr(ezdxf.enums.TextEntityAlignment, align))
        self._reg(self._tbox(txt, at, h, rot, align))

    # --- keeping things apart --------------------------------------------------------
    CELL = 2000.0

    def _cells(self, g):
        x0, y0, x1, y1 = g.bounds
        c = self.CELL
        for i in range(int(x0 // c), int(x1 // c) + 1):
            for j in range(int(y0 // c), int(y1 // c) + 1):
                yield i, j

    def _reg(self, g):
        """Mark an area of the plan as taken."""
        if g is None or g.is_empty:
            return
        for k in self._cells(g):
            self.occ[k].append(g)

    def _overlap(self, g):
        """How much of g is already taken (to pick the least bad place when none is free)."""
        g = g.buffer(0.15 * self.th, join_style=2)
        seen, tot = set(), 0.0
        for k in self._cells(g):
            for o in self.occ.get(k, ()):
                if id(o) not in seen:
                    seen.add(id(o))
                    if o.intersects(g):
                        tot += o.intersection(g).area
        return tot

    def _clear(self, g, margin=None):
        g = g.buffer(0.15 * self.th if margin is None else margin, join_style=2)
        seen = set()
        for k in self._cells(g):
            for o in self.occ.get(k, ()):
                if id(o) not in seen:
                    seen.add(id(o))
                    if o.intersects(g):
                        return False
        return True

    @staticmethod
    def _tbox(txt, at, h, rot=0, align="MIDDLE_CENTER"):
        """Area a text takes (Arial-like proportions)."""
        w = 0.8 * h * len(txt)
        v, hz = align.split("_")
        x0 = {"LEFT": 0.0, "CENTER": -w / 2, "RIGHT": -w}[hz]
        y0 = {"BOTTOM": 0.0, "MIDDLE": -h / 2, "TOP": -h}[v]
        c, s_ = math.cos(math.radians(rot)), math.sin(math.radians(rot))
        return Polygon([(at[0] + x * c - y * s_, at[1] + x * s_ + y * c)
                        for x, y in ((x0, y0), (x0 + w, y0), (x0 + w, y0 + h), (x0, y0 + h))])

    def dim_auto(self, a, b, at, line, horiz, txt, pref=0):
        """Dimension a..b along X (horiz) or Y, measured at 'at', dimension line at 'line'. The text goes
        to the first free place: between the extension lines (if it fits), beside the chain end
        (pref +1 after it, -1 before it), on the other side of the line, or a row further out."""
        th = self.th
        tw = 0.8 * 1.2 * th * len(txt)
        r = 0.9 * th                                  # text centre off the dimension line
        up = 1 if horiz else -1                       # text above X dimensions, left of Y dimensions
        mid = (a + b) / 2
        fits = b - a >= tw + 0.4 * th
        before, after = a - 0.4 * th - tw / 2, b + 0.4 * th + tw / 2
        ends = [after, before] if pref >= 0 else [before, after]
        cands = []
        if fits:                                      # along the segment first, both sides of the line
            room = (b - a - tw) / 2 - 0.2 * th
            for f in (0, -0.4, 0.4, -0.8, 0.8):
                cands += [(mid + f * room, up * r), (mid + f * room, -up * r)]
            if pref != 0:
                cands[2:2] = [(ends[0], up * r)]
        else:
            cands += [(ends[0], up * r), (ends[1], up * r), (ends[0], -up * r), (ends[1], -up * r)]
        cands += [(mid, up * (r + 1.5 * th)), (mid, -up * (r + 1.5 * th)),
                  (ends[0], up * (r + 1.5 * th)), (ends[1], up * (r + 1.5 * th))]
        if fits:
            cands += [(ends[0], up * r), (ends[1], up * r)]
        rot = 0 if horiz else 90
        pick = None
        for sa, off in cands:
            c = (sa, line + off) if horiz else (line + off, sa)
            bx = self._tbox(txt, c, 1.2 * th, rot)
            if self._clear(bx):
                pick = (c, bx)
                break
        if pick is None:                              # nowhere free: the least covered place
            opts = []
            for sa, off in cands:
                c = (sa, line + off) if horiz else (line + off, sa)
                bx = self._tbox(txt, c, 1.2 * th, rot)
                opts.append((self._overlap(bx), c, bx))
            _, c, bx = min(opts, key=lambda o: o[0])
            pick = (c, bx)
        self._reg(pick[1])
        if horiz:
            self.dim((a, at), (b, at), (a, line), 0, text_at=pick[0], text=txt)
        else:
            self.dim((at, a), (at, b), (line, a), 90, text_at=pick[0], text=txt)

    def _fmt(self, length, half=True):
        unit = UNIT_MM.get(self.cfg.get("drawing", {}).get("dimension_unit", "cm"), 10.0)
        v = round(length / unit * 2) / 2 if half else round(length / unit)
        return f"{v:.0f}" if v == int(v) else f"{v:.1f}"

    def dim(self, p1, p2, base, angle, text_at=None, text="<>"):
        try:
            d = self.msp.add_linear_dim(base=self.m(base), p1=self.m(p1), p2=self.m(p2), angle=angle,
                                        text=text, dimstyle="FW_DIM", dxfattribs={"layer": "FW_DIM"})
            if text_at is not None:
                d.set_location(self.m(text_at), leader=False, relative=False)
            d.render()
            if d.dimension.dxf.get("text_midpoint") is None:     # a blank text leaves none (readers need it)
                d.dimension.dxf.text_midpoint = self.m(text_at if text_at is not None else base)
        except Exception as ex:
            warn(f"formwork dimension could not be drawn: {ex}")

    def hatch(self, poly, layer, pattern, scale, color=256, solid_bg=None):
        polys = getattr(poly, "geoms", [poly])
        for pg in polys:
            if pg.is_empty or pg.geom_type != "Polygon" or pg.area < 1:
                continue
            if solid_bg is not None:
                hb = self.msp.add_hatch(color=solid_bg, dxfattribs={"layer": layer})
                hb.paths.add_polyline_path([self.m(c) for c in pg.exterior.coords], is_closed=True, flags=1)
                for r in pg.interiors:
                    hb.paths.add_polyline_path([self.m(c) for c in r.coords], is_closed=True, flags=0)
            h = self.msp.add_hatch(color=color, dxfattribs={"layer": layer})
            h.set_pattern_fill(pattern, scale=scale)
            h.paths.add_polyline_path([self.m(c) for c in pg.exterior.coords], is_closed=True, flags=1)
            for r in pg.interiors:
                h.paths.add_polyline_path([self.m(c) for c in r.coords], is_closed=True, flags=0)

    # --- parts ----------------------------------------------------------------
    def hatch_scale(self):
        return 0.8 * self.th / self.u / 7.5        # ~3 cm pattern spacing at text height 15 cm

    def hatch_elements(self, cut=None):
        """Cross-hatch of the shear walls and columns (also used on the reinforcement plans).
        'cut' (the wall sections) is left out of the wall hatch, so the section sits on clean ground."""
        walls = self.G.walls_union
        if cut is not None and not cut.is_empty and not walls.is_empty:
            walls = walls.difference(cut)
        if not walls.is_empty:
            self.hatch(walls, "FW_WALL_HATCH", "ANSI37", self.hatch_scale())
        for c in self.G.columns:
            self.hatch(c.poly, "FW_COLUMN_HATCH", "ANSI37", self.hatch_scale())

    def labels(self):
        """Names of the walls (SWn X x Y) and columns, each at the first free place around the element."""
        G, th = self.G, self.th
        order = sorted(G.walls, key=lambda w: (-round(w.cy / 500), w.cx))
        sec = {id(w): tc for w, tc, _ in self.sections}
        k = 0
        for w in order:
            if w.orient is None or w.length < 3 * w.thk or self._is_core_wall(w):
                continue
            k += 1
            x0, y0, x1, y1 = w.poly.bounds
            txt = f"SW{k} {(x1 - x0) / 10:.0f}X{(y1 - y0) / 10:.0f}"
            tw = 0.7 * th * len(txt)
            ts = [(w.lo + w.hi) / 2]
            if id(w) in sec:                       # beside the wall, clear of its section
                tc = sec[id(w)]
                half = self.P.h / 2 + self.ws_stub
                a, b = (w.lo, tc - half), (tc + half, w.hi)
                ts = sorted([sum(a) / 2, sum(b) / 2], key=lambda t: -(b[1] - b[0] if t > tc else a[1] - a[0]))
            ts += [w.lo + tw / 2, w.hi - tw / 2, (w.lo + w.hi) / 2]
            if w.orient == "x":
                inside = 1 if self._inside((w.cx, w.cy + w.thk / 2 + 1.2 * th)) else -1
                cands = [((t, w.cy + sd * (w.thk / 2 + row * th)), 0, "BOTTOM_CENTER" if sd > 0 else "TOP_CENTER")
                         for row in (0.9, 2.6) for sd in (inside, -inside) for t in ts]
            else:
                inside = 1 if self._inside((w.cx + w.thk / 2 + 1.2 * th, w.cy)) else -1
                cands = [((w.cx + sd * (w.thk / 2 + row * th), t), 90, "TOP_CENTER" if sd > 0 else "BOTTOM_CENTER")
                         for row in (0.9, 2.6) for sd in (inside, -inside) for t in ts]
            self._place(txt, th, cands)
        # columns: name  Cn-b/h  or  Cn-Ød, below the column if free, else around it
        cols = sorted(G.columns, key=lambda c: (-round(c.cy / 500), c.cx))
        for i, c in enumerate(cols, 1):
            size = f"Ø{c.wx / 10:.0f}" if c.round else f"{c.wx / 10:.0f}/{c.wy / 10:.0f}"
            x0, y0, x1, y1 = c.poly.bounds
            g = 0.5 * th
            cands = [((c.cx, y0 - g), 0, "TOP_CENTER"), ((c.cx, y1 + g), 0, "BOTTOM_CENTER"),
                     ((x1 + g, c.cy), 0, "MIDDLE_LEFT"), ((x0 - g, c.cy), 0, "MIDDLE_RIGHT"),
                     ((c.cx, y0 - 2.2 * th), 0, "TOP_CENTER"), ((c.cx, y1 + 2.2 * th), 0, "BOTTOM_CENTER"),
                     ((x1 + g, y0 - g), 0, "TOP_LEFT"), ((x0 - g, y0 - g), 0, "TOP_RIGHT")]
            self._place(f"C{i}-{size}", 0.9 * th, cands)

    def _place(self, txt, h, cands):
        """Write a text at the first candidate place (at, rotation, alignment) that is free."""
        for at, rot, align in cands:
            if self._clear(self._tbox(txt, at, h, rot, align)):
                break
        else:                                         # nowhere free: the least covered place
            at, rot, align = min(cands, key=lambda c: self._overlap(self._tbox(txt, c[0], h, c[1], c[2])))
        self.text(txt, at, h, rot=rot, align=align)

    # --- sections through the shear walls ---------------------------------------
    def _core_holes(self):
        """Cores: the inside of a closed box of walls, or an opening walled on most of its perimeter."""
        if self._cores is None:
            G = self.G
            boxes = [Polygon(r) for g in getattr(G.walls_union, "geoms", [G.walls_union])
                     if not g.is_empty for r in g.interiors]
            near = G.walls_union.buffer(30) if not G.walls_union.is_empty else Polygon()
            walled = [hole for hole in G.holes if not near.is_empty and
                      hole.exterior.intersection(near).length >= 0.5 * hole.exterior.length]
            self._cores = boxes + [hp for hp in walled if not any(b.buffer(10).contains(hp) for b in boxes)]
        return self._cores

    def _is_core_wall(self, w):
        return w.core or any(w.poly.distance(hp) < 30 for hp in self._core_holes())

    @staticmethod
    def _wave(s_at, r0, r1, amp, over=0.0, n=16):
        """Break line across a bar end: one sine wave from r0 to r1 (optionally running on a bit)."""
        out = []
        lo, hi = -over, 1 + over
        for k in range(n + 1):
            f = lo + (hi - lo) * k / n
            out.append((s_at + amp * math.sin(2 * math.pi * f), r0 + f * (r1 - r0)))
        return out

    def _bar(self, s0, s1, r0, r1, wave0, wave1):
        """Outline (s, r) of a bar from s0 to s1, width r0..r1, with wavy (broken) ends where asked."""
        amp = 0.15 * (r1 - r0)
        e1 = self._wave(s1, r0, r1, amp) if wave1 else [(s1, r0), (s1, r1)]
        e0 = self._wave(s0, r0, r1, amp)[::-1] if wave0 else [(s0, r1), (s0, r0)]
        return e1 + e0

    def wall_sections(self):
        """A section through every shear wall, laid flat on the wall as on the drawings: the slab
        (grey, as thick as the slab) runs across the wall, the wall (grey, as thick as the wall) shows
        above and below the slab, all ends broken off with wavy lines. Slab on both sides = cross,
        slab on one side only (edge walls, core walls) = T. With the slab thickness, the level of the
        slab on its top face, and the wall thickness from the axis. Returns the area covered."""
        G, P, th = self.G, self.P, self.th
        ws = self.fw.get("wall_sections") or {}
        if not ws.get("enabled", True):
            return Polygon()
        h = P.h
        ext = float(ws.get("slab_extension", 450))         # slab shown beyond each wall face
        stub = self.ws_stub
        cols = [c.poly for c in G.columns]
        lvl = self.fw.get("section_level") or self.fw.get("level") or ""
        lvl = (lvl[0] + " " + lvl[1:].lstrip()) if lvl and lvl[0] in "+-±" else lvl
        cover = []
        for w in G.walls:
            if w.orient is None or w.length < h + 2 * stub:
                continue
            others = unary_union([x.poly for x in G.walls if x is not w] + cols)
            n0, n1 = (w.cx - w.thk / 2, w.cx + w.thk / 2) if w.orient == "y" else \
                     (w.cy - w.thk / 2, w.cy + w.thk / 2)
            xy = (lambda s_, r_: (s_, r_)) if w.orient == "y" else (lambda s_, r_: (r_, s_))   # (n, t) -> (x, y)
            L = w.length
            best = None
            for tc in (w.lo + L / 3, w.lo + 2 * L / 3, w.lo + L / 2):
                tc = min(max(tc, w.lo + h / 2 + stub), w.hi - h / 2 - stub)
                sides = []
                for sd in (-1, 1):
                    face = n0 if sd < 0 else n1
                    probe = Point(*xy(face + sd * min(200, ext / 2), tc))
                    if not G.slab.contains(probe) or others.contains(probe):
                        continue
                    for e in (ext, 0.6 * ext):
                        a, b = sorted((face, face + sd * e))
                        rect = Polygon([xy(a, tc - h / 2 - 50), xy(b, tc - h / 2 - 50),
                                        xy(b, tc + h / 2 + 50), xy(a, tc + h / 2 + 50)])
                        if G.slab.buffer(5).contains(rect) and not rect.intersects(others.buffer(-1)):
                            sides.append((sd, e))
                            break
                if sides and (best is None or len(sides) > len(best[1])):
                    best = (tc, sides)
                if best and len(best[1]) == 2:
                    break
            if best is None:
                continue
            tc, sides = best
            self.sections.append((w, tc, sides))
            # slab strip (s = n across the wall, r = t along it) and wall stub (s = t, r = n)
            e_lo = dict(sides).get(-1)
            e_hi = dict(sides).get(1)
            s0 = n0 - e_lo if e_lo else (n0 + n1) / 2
            s1 = n1 + e_hi if e_hi else (n0 + n1) / 2
            strip = [xy(a, b) for a, b in self._bar(s0, s1, tc - h / 2, tc + h / 2, bool(e_lo), bool(e_hi))]
            t0, t1 = tc - h / 2 - stub, tc + h / 2 + stub
            wave0 = wave1 = True
            if is_mat():                    # the mat is the lowest slab: the wall only stands on it
                top_hi = w.orient == "y"    # top face of the section: +t for walls along Y, -t along X
                if top_hi:
                    t0, wave0 = tc - h / 2, False
                else:
                    t1, wave1 = tc + h / 2, False
            wall = [xy(b, a) for a, b in self._bar(t0, t1, n0, n1, wave0, wave1)]
            shape = unary_union([Polygon(strip).buffer(0), Polygon(wall).buffer(0)])
            cover.append(shape)
            self._reg(shape)
            for pg in getattr(shape, "geoms", [shape]):
                hb = self.msp.add_hatch(color=8, dxfattribs={"layer": "FW_WALL_SECTION"})
                hb.set_solid_fill(color=8)
                hb.paths.add_polyline_path([self.m(c) for c in pg.exterior.coords], is_closed=True)
                self.msp.add_lwpolyline([self.m(c) for c in pg.exterior.coords], close=True,
                                        dxfattribs={"layer": "FW_WALL_SECTION", "lineweight": 50})
            # break lines run on a little past the outline, as drawn by hand
            amp_s, amp_w = 0.15 * h, 0.15 * w.thk
            ends = [(s_, tc - h / 2, tc + h / 2, amp_s, False) for s_, e_ in ((s0, e_lo), (s1, e_hi)) if e_]
            ends += [(t_, n0, n1, amp_w, True) for t_, wv in ((t0, wave0), (t1, wave1)) if wv]
            for s_, r0, r1, amp, flip in ends:
                pts = self._wave(s_, r0, r1, amp, over=0.15)
                pts = [xy(b, a) if flip else xy(a, b) for a, b in pts]
                self.msp.add_lwpolyline([self.m(c) for c in pts],
                                        dxfattribs={"layer": "FW_WALL_SECTION", "lineweight": 25})
            # slab thickness, just beyond the broken end of the slab on the first slab side
            sd, e = sides[0]
            face = n0 if sd < 0 else n1
            end = face + sd * (e + 1.1 * th)
            p1, p2 = xy(face, tc - h / 2), xy(face, tc + h / 2)
            tat = xy(end + sd * 0.9 * th, tc)                    # text on the far side of the line
            self.dim(p1, p2, xy(end, tc - h / 2), 90 if w.orient == "y" else 0, text_at=tat, text=self._fmt(h))
            self._reg(self._tbox(self._fmt(h), tat, 1.2 * th, 90 if w.orient == "y" else 0))
            # level of the slab: triangle on the top face, leader and level text
            if lvl:
                top, out = (tc + h / 2, 1) if w.orient == "y" else (tc - h / 2, -1)
                na, tri = face + sd * 0.45 * e, 0.57 * th
                self.level_mark(xy, na, top, out, -sd, tri, lvl, w.orient)
        return unary_union(cover) if cover else Polygon()

    def level_mark(self, xy, na, top, out, to_wall, tri, txt, orient):
        """Level symbol: triangle with its point on the top face of the slab (half filled, the filled
        half towards the wall), a leader out of the slab and along it, the level written on it."""
        th = self.th
        apex, base = xy(na, top), top + out * tri
        left, right = xy(na - tri, base), xy(na + tri, base)
        mid = xy(na, base)
        lay = {"layer": "FW_WALL_SECTION"}
        self.msp.add_lwpolyline([self.m(apex), self.m(left), self.m(right)], close=True, dxfattribs=lay)
        half = [apex, xy(na + to_wall * tri, base), mid]
        hb = self.msp.add_hatch(color=7, dxfattribs=lay)
        hb.set_solid_fill(color=7)
        hb.paths.add_polyline_path([self.m(c) for c in half], is_closed=True)
        far = na - to_wall * 5 * th
        knee = top + out * th
        self.msp.add_lwpolyline([self.m(apex), self.m(xy(na, knee)), self.m(xy(far, knee))], dxfattribs=lay)
        self._reg(LineString([apex, xy(na, knee), xy(far, knee)]).buffer(0.2 * th))
        # text on the leader, starting at its far end and reading towards the wall
        at = xy(far, knee + out * 0.25 * th)
        self.text(txt, at, 4 / 3 * th, rot=0 if orient == "y" else 90,
                  align="BOTTOM_LEFT" if to_wall > 0 else "BOTTOM_RIGHT",
                  layer="FW_WALL_SECTION")

    # --- dimensions of every element ----------------------------------------------
    def dim_chain(self, pts, at, horiz, sg):
        """Dimension chain through the points pts (along X if horiz, else along Y), with its line
        0.8 x text height beyond 'at' on side sg. Values to 0.5 cm (12.5). Each text goes to a free
        place: the first one before the chain, the last one after it if they do not fit."""
        line = at + sg * 0.8 * self.th
        segs = [(a, b) for a, b in zip(pts, pts[1:]) if b - a > 1]
        for i, (a, b) in enumerate(segs):
            pref = 0 if len(segs) == 1 else (-1 if i == 0 else (1 if i == len(segs) - 1 else 0))
            self.dim_auto(a, b, at, line, horiz, self._fmt(b - a), pref=pref)
        if segs:                                      # the chain's line is taken too
            q = 0.25 * self.th
            self._reg(box(pts[0], line - q, pts[-1], line + q) if horiz else box(line - q, pts[0], line + q, pts[-1]))
        return len(segs)

    def _band(self, lo, hi, at, horiz, sg):
        """Area a dimension chain on side sg would take."""
        th = self.th
        a, b = at + sg * 0.5 * th, at + sg * 1.9 * th
        a, b = min(a, b), max(a, b)
        return box(lo, a, hi, b) if horiz else box(a, lo, b, hi)

    def _side(self, lo, hi, at_lo, at_hi, horiz, prefer):
        """Side (+1/-1) for a chain from lo to hi beside an element spanning at_lo..at_hi across:
        the preferred side if free (inside the slab, nothing in the way), else the other one."""
        th = self.th
        for sg in (prefer, -prefer):
            at = at_hi if sg > 0 else at_lo
            mid = (lo + hi) / 2
            probe = (mid, at + sg * 1.2 * th) if horiz else (at + sg * 1.2 * th, mid)
            if self._inside(probe) and self._clear(self._band(lo, hi, at, horiz, sg)):
                return sg, at
        sg = prefer
        return sg, (at_hi if sg > 0 else at_lo)

    def element_dimensions(self):
        """Every wall, column and opening dimensioned on its own: wall thickness and column sizes split
        at the axes running through them (face - axis - face), wall lengths where no grid line runs
        along the wall, opening sizes."""
        G, th = self.G, self.th
        xs, ys, _ = self.grid()
        sec = {id(w): tc for w, tc, _ in self.sections}
        holes = self._core_holes()

        def chain(lo, hi, axes):
            return [lo] + [a for a in sorted(axes) if lo + 20 < a < hi - 20] + [hi]

        def free(p):
            q = Point(*p)
            return self._inside(p) and not G.walls_union.contains(q) and \
                not any(c.poly.contains(q) for c in G.columns)
        n = 0
        for w in G.walls:
            if w.orient is None:
                continue
            x0, y0, x1, y1 = w.poly.bounds
            # thickness, at the wall end next to its section (or the other end if that one is not free)
            if w.orient == "y":
                ends = [(y0, -1), (y1, 1)]
                if id(w) in sec and sec[id(w)] > (y0 + y1) / 2:
                    ends.reverse()
                ends = [e for e in ends if free((w.cx, e[0] + e[1] * 1.2 * th))]   # not at a junction
                ends.sort(key=lambda e: not self._clear(self._band(x0, x1, e[0], True, e[1])))
                if ends:
                    n += self.dim_chain(chain(x0, x1, xs), ends[0][0], True, ends[0][1])
            else:
                ends = [(x0, -1), (x1, 1)]
                if id(w) in sec and sec[id(w)] > (x0 + x1) / 2:
                    ends.reverse()
                ends = [e for e in ends if free((e[0] + e[1] * 1.2 * th, w.cy))]   # not at a junction
                ends.sort(key=lambda e: not self._clear(self._band(y0, y1, e[0], False, e[1])))
                if ends:
                    n += self.dim_chain(chain(y0, y1, ys), ends[0][0], False, ends[0][1])
            # length, where no grid line runs along the wall (the grid chains already give it there)
            along = xs if w.orient == "y" else ys
            lo_, hi_ = (x0, x1) if w.orient == "y" else (y0, y1)
            if w.length >= 3 * w.thk and not any(lo_ - 1 <= a <= hi_ + 1 for a in along):
                if w.orient == "x":
                    sg, at = self._side(x0, x1, y0, y1, True, -1)
                    n += self.dim_chain([x0, x1], at, True, sg)
                else:
                    sg, at = self._side(y0, y1, x0, x1, False, 1)
                    n += self.dim_chain([y0, y1], at, False, sg)
        for c in G.columns:
            x0, y0, x1, y1 = c.poly.bounds
            if not c.round and c.poly.area < 0.9 * (x1 - x0) * (y1 - y0):
                continue                                  # turned column: its sizes are in its name
            sg, at = self._side(x0, x1, y0, y1, True, 1)  # across X, above (or below) the column
            n += self.dim_chain(chain(x0, x1, xs), at, True, sg)
            sg, at = self._side(y0, y1, x0, x1, False, 1)  # across Y, right (or left) of the column
            n += self.dim_chain(chain(y0, y1, ys), at, False, sg)
        for hole in G.holes:
            if any(hp.buffer(10).contains(hole) for hp in holes):
                continue                                  # the void inside a core: walls dimension it
            x0, y0, x1, y1 = hole.bounds
            n += self.dim_chain([x0, x1], y0, True, 1)       # inside the opening, along two sides
            n += self.dim_chain([y0, y1], x0, False, 1)
        return n

    def _inside(self, p):
        return self.G.slab.buffer(-10).contains(Point(*p))

    def openings(self):
        for hole in self.G.holes:
            minx, miny, maxx, maxy = hole.bounds
            for a, b in (((minx, miny), (maxx, maxy)), ((minx, maxy), (maxx, miny))):
                ln = LineString([a, b]).intersection(hole)
                for g in getattr(ln, "geoms", [ln]):
                    if g.geom_type == "LineString" and g.length > 0:
                        self.msp.add_line(self.m(g.coords[0]), self.m(g.coords[-1]),
                                          dxfattribs={"layer": "FW_OPENING"})
            core = any(hp.buffer(10).contains(hole) for hp in self._core_holes())
            self.text("CORE" if core else "OPENING", hole.representative_point().coords[0],
                      self.th if core else 0.9 * self.th)
        for hp in self._core_holes():                 # a closed box of walls without an opening drawn
            if not any(hp.buffer(10).contains(hole) for hole in self.G.holes):
                self.text("CORE", hp.representative_point().coords[0], self.th)

    def grid(self):
        """Axis positions: drawn axes if present, otherwise the column lines."""
        ax = self.G.axes
        if ax["x"] and ax["y"]:
            return [a[0] for a in ax["x"]], [a[0] for a in ax["y"]], True
        # no axes drawn: use the centre lines of the columns and walls
        return list(self.G.sup_lines["x"]), list(self.G.sup_lines["y"]), False

    def slab_sections(self):
        """Rotated slab sections (grey strip, width = slab thickness) at mid-span on every grid line,
        each with its thickness dimension, as in a usual formwork plan."""
        G, P, th = self.G, self.P, self.th
        h = P.h
        xs, ys, _ = self.grid()
        solid = G.slab.buffer(-100)
        blocked = unary_union([G.walls_union.buffer(150)] + [c.poly.buffer(150) for c in G.columns] +
                              [self.ws_cover.buffer(300 + 3 * th)])     # incl. its thickness dimension
        placed = 0
        self.strip_half = defaultdict(float)     # (grid line dirn, coordinate) -> half length of its strips
        self.strip_at = {}                       # (grid line dirn, coordinate) -> positions of its strips
        for dirn, lines, others in (("v", ys, xs), ("h", xs, ys)):
            # dirn v: strip long in Y, sitting on a horizontal grid line, between two vertical axes
            for gl in lines:
                for p, q in zip(others, others[1:]):
                    span = q - p
                    if span < 1500:
                        continue
                    mid = (p + q) / 2
                    # not longer than the room to the next parallel grid line (no strips running together)
                    i = lines.index(gl)
                    room = min([abs(gl - v) for v in lines[max(i - 1, 0):i + 2] if v != gl] or [4000])
                    for frac in (0.35, 0.25, 0.18):
                        Ls = min(max(frac * span, 800, 2 * h), 2000, 0.7 * room)
                        if Ls < 600:
                            break
                        if dirn == "v":
                            strip = box(mid - h / 2, gl - Ls / 2, mid + h / 2, gl + Ls / 2)
                        else:
                            strip = box(gl - Ls / 2, mid - h / 2, gl + Ls / 2, mid + h / 2)
                        if solid.contains(strip) and not strip.intersects(blocked) and \
                                self._clear(strip.buffer(1.5 * th)):
                            self._section(strip, dirn)
                            self.strip_half[(dirn, round(gl))] = max(self.strip_half[(dirn, round(gl))], Ls / 2)
                            self.strip_at.setdefault((dirn, round(gl)), []).append(mid)
                            placed += 1
                            break
        return placed

    def _section(self, strip, dirn):
        th = self.th
        x0, y0, x1, y1 = strip.bounds
        hb = self.msp.add_hatch(color=8, dxfattribs={"layer": "FW_SLAB_SECTION"})
        hb.set_solid_fill(color=8)
        hb.paths.add_polyline_path([self.m(c) for c in strip.exterior.coords], is_closed=True)
        lw = {"layer": "FW_SLAB_SECTION", "lineweight": 50, "color": 7}
        txt = self._fmt(self.P.h)
        if dirn == "v":
            self.msp.add_line(self.m((x0, y0)), self.m((x0, y1)), dxfattribs=lw)
            self.msp.add_line(self.m((x1, y0)), self.m((x1, y1)), dxfattribs=lw)
            tat = (x1 + 1.6 * th, y0 - 1.0 * th + 0.8 * th)
            self.dim((x0, y0), (x1, y0), (x0, y0 - 1.0 * th), 0, text_at=tat, text=txt)
            self._reg(self._tbox(txt, tat, 1.2 * th))
        else:
            self.msp.add_line(self.m((x0, y0)), self.m((x1, y0)), dxfattribs=lw)
            self.msp.add_line(self.m((x0, y1)), self.m((x1, y1)), dxfattribs=lw)
            tat = (x1 + 1.0 * th - 0.8 * th, y1 + 1.6 * th)
            self.dim((x1, y0), (x1, y1), (x1 + 1.0 * th, y0), 90, text_at=tat, text=txt)
            self._reg(self._tbox(txt, tat, 1.2 * th, 90))
        self._reg(strip.buffer(0.3 * th))

    def internal_dimensions(self, quiet_columns=False):
        """Dimension chains inside the slab along every grid line, through the columns and shear walls:
        slab edge - element face - element width - clear span - ... - slab edge (as in a usual formwork plan).
        The dimension line lies on the grid line itself, so it passes through the elements; where a
        slab-section symbol sits on the line, the text is moved beside it."""
        G = self.G
        xs, ys, _ = self.grid()
        elems = unary_union([G.walls_union] + [c.poly for c in G.columns]) if G.columns or \
            not G.walls_union.is_empty else Polygon()
        minx, miny, maxx, maxy = G.bounds
        n = 0
        for dirn, lines in (("v", ys), ("h", xs)):
            # dirn v: horizontal grid line y = gl, chain along X drawn below it
            # dirn h: vertical grid line x = gl, chain along Y drawn left of it
            for gl in lines:
                if dirn == "v":
                    ln = LineString([(minx - 10, gl), (maxx + 10, gl)])
                else:
                    ln = LineString([(gl, miny - 10), (gl, maxy + 10)])
                pts = []
                for geom in (G.slab, elems):
                    if geom.is_empty:
                        continue
                    g = ln.intersection(geom)
                    for part in getattr(g, "geoms", [g]):
                        if part.geom_type == "LineString" and part.length > 1:
                            for c in part.coords:
                                pts.append(c[0] if dirn == "v" else c[1])
                if len(pts) < 2:
                    continue
                pts.sort()
                uniq = [pts[0]]
                for p in pts[1:]:
                    if p - uniq[-1] > 20:
                        uniq.append(p)
                # keep only segments lying on the slab or inside an element
                segs = []
                for a, b in zip(uniq, uniq[1:]):
                    mid = (a + b) / 2
                    mp = Point(mid, gl) if dirn == "v" else Point(gl, mid)
                    if G.slab.buffer(1).contains(mp) or elems.buffer(1).contains(mp):
                        segs.append((a, b))
                if len(segs) < 2:              # nothing crossed on this line
                    continue
                for a, b in segs:                 # texts at free places (sections, names, other texts)
                    mp = Point((a + b) / 2, gl) if dirn == "v" else Point(gl, (a + b) / 2)
                    if quiet_columns and any(c.poly.contains(mp) for c in G.columns):
                        # across a column: the column's own chain gives the value, keep the line only
                        mid = (a + b) / 2
                        if dirn == "v":
                            self.dim((a, gl), (b, gl), (a, gl), 0, text_at=(mid, gl), text=" ")
                        else:
                            self.dim((gl, a), (gl, b), (gl, a), 90, text_at=(gl, mid), text=" ")
                    else:
                        self.dim_auto(a, b, gl, gl, dirn == "v", self._fmt(b - a, half=False))
                    n += 1
        return n

    def axes_and_dimensions(self, with_dims=True):
        """Axis bubbles (and on the formwork plan also the outer dimension chains)."""
        G, th = self.G, self.th
        minx, miny, maxx, maxy = G.bounds
        xs, ys, drawn = self.grid()
        r = 2.2 * th
        d1, d2 = 4 * th, 8 * th           # distance of the two dimension rows from the slab
        # dimension chains: slab edge - axes - slab edge, and the overall size
        chx = sorted(set([round(minx)] + [round(x) for x in xs if minx - 1 <= x <= maxx + 1] + [round(maxx)]))
        chy = sorted(set([round(miny)] + [round(y) for y in ys if miny - 1 <= y <= maxy + 1] + [round(maxy)]))
        for yb, side in (((maxy + d1, 1), (miny - d1, -1)) if with_dims else ()):
            for a, b in zip(chx, chx[1:]):
                if b - a > 5:
                    self.dim((a, maxy if side > 0 else miny), (b, maxy if side > 0 else miny), (a, yb), 0)
        if with_dims:
            self.dim((minx, maxy), (maxx, maxy), (minx, maxy + d2), 0)
        for xb, side in (((minx - d1, -1), (maxx + d1, 1)) if with_dims else ()):
            for a, b in zip(chy, chy[1:]):
                if b - a > 5:
                    self.dim((minx if side < 0 else maxx, a), (minx if side < 0 else maxx, b), (xb, a), 90)
        if with_dims:
            self.dim((minx, miny), (minx, maxy), (minx - d2, miny), 90)
        # axis bubbles: 1, 2, 3 ... (vertical axes, left to right), A, B, C ... (horizontal, bottom to top)
        reach = d2 + 3 * th + r
        letters = [chr(ord("A") + i) for i in range(26)]
        axx = self.G.axes["x"] if drawn else [(x, miny, maxy) for x in xs]
        axy = self.G.axes["y"] if drawn else [(y, minx, maxx) for y in ys]
        if not drawn:                              # no axis lines in the drawing: draw them
            for x, lo, hi in axx:
                self.msp.add_line(self.m((x, lo)), self.m((x, hi)), dxfattribs={"layer": "FW_AXIS", "color": 8})
            for y, lo, hi in axy:
                self.msp.add_line(self.m((lo, y)), self.m((hi, y)), dxfattribs={"layer": "FW_AXIS", "color": 8})
        for i, (x, lo, hi) in enumerate(axx):
            for end, sgn in ((max(hi, maxy + reach - r), 1), (min(lo, miny - reach + r), -1)):
                self._bubble((x, end + sgn * r), str(i + 1), r, (x, hi if sgn > 0 else lo), (x, end))
        for i, (y, lo, hi) in enumerate(axy):
            lab = letters[i] if i < 26 else f"A{i - 25}"
            for end, sgn in ((min(lo, minx - reach + r), -1), (max(hi, maxx + reach - r), 1)):
                self._bubble((end + sgn * r, y), lab, r, (lo if sgn < 0 else hi, y), (end, y))

    def _bubble(self, c, lab, r, ax_end, new_end):
        if math.dist(ax_end, new_end) > 1:          # extend the axis to its bubble
            self.msp.add_line(self.m(ax_end), self.m(new_end), dxfattribs={"layer": "FW_AXIS", "linetype": "DASHDOT"
                                                                            if "DASHDOT" in self.doc.linetypes else "CONTINUOUS"})
        self.msp.add_circle(self.m(c), r / self.u, dxfattribs={"layer": "FW_AXIS"})
        self.text(lab, c, 1.6 * self.th, layer="FW_AXIS")

    def outer_extents(self):
        """Top and right edge of everything drawn around the plan (dimensions, axes, bubbles)."""
        G, th = self.G, self.th
        minx, miny, maxx, maxy = G.bounds
        r = 2.2 * th
        reach = 8 * th + 3 * th + r
        top = max([maxy + reach] + [hi + r for _, _, hi in G.axes["x"]]) + r
        right = max([maxx + reach] + [hi + r for _, _, hi in G.axes["y"]]) + r
        return top, right

    def legend_and_notes(self, title):
        G, P, th = self.G, self.P, self.th
        minx, miny, maxx, maxy = G.bounds
        fw = self.fw
        top, right = self.outer_extents()
        lvl = fw.get("level")
        ttl = title + (f" {lvl}" if lvl else "")
        self.text(ttl, ((minx + maxx) / 2, top + 5 * th), 2.5 * th)
        self.text(f"M 1:{fw.get('scale', 50)}", ((minx + maxx) / 2, top + 2.5 * th), th)
        # legend + notes to the right of the plan
        x = right + 4 * th
        y = maxy
        self.text("LEGEND", (x, y), 1.4 * th, align="MIDDLE_LEFT")
        items = [("FW_WALL_HATCH", "ANSI37", "Shear wall"), ("FW_COLUMN_HATCH", "ANSI37", "Column"),
                 ("FW_SLAB_SECTION", "SOLID", f"R.C. {'mat' if is_mat() else 'slab'} section (h = {P.h / 10:.0f} cm)")]
        for k, (layer, pat, lab) in enumerate(items):
            yy = y - (3 + 2.5 * k) * th
            sw = box(x, yy - 0.6 * th, x + 5 * th, yy + 0.6 * th)
            if pat == "SOLID":
                hb = self.msp.add_hatch(color=8, dxfattribs={"layer": layer})
                hb.set_solid_fill(color=8)
                hb.paths.add_polyline_path([self.m(c) for c in sw.exterior.coords], is_closed=True)
            else:
                self.hatch(sw, layer, pat, self.hatch_scale())
            self.msp.add_lwpolyline([self.m(c) for c in sw.exterior.coords], close=True,
                                    dxfattribs={"layer": "FW_TEXT"})
            self.text(lab, (x + 6 * th, yy), th, align="MIDDLE_LEFT")
        notes = ["NOTES:",
                 f"{'Mat foundation' if is_mat() else 'Slab'} thickness h = {P.h / 10:.0f} cm (everywhere)",
                 f"Concrete C{P.fck:.0f}/{ {12: 15, 16: 20, 20: 25, 25: 30, 30: 37, 35: 45}.get(int(P.fck), '')}",
                 f"Steel B{P.fyk:.0f}B",
                 f"Cover: bottom {P.cb / 10:.0f} cm, top {P.ct / 10:.0f} cm"]
        notes += [str(n) for n in (fw.get("notes") or [])]
        for k, n in enumerate(notes):
            self.text(n, (x, y - (12 + 1.8 * k) * th), (1.2 if k == 0 else 1.0) * th, align="MIDDLE_LEFT")

    def draw(self):
        # first what has a fixed place, then the names, then the dimensions (their texts move to free places)
        self.ws_cover = self.wall_sections()
        self.hatch_elements(cut=self.ws_cover)
        self.openings()
        n = self.slab_sections()
        if n == 0:
            warn("formwork plan: no free place found for the slab section symbols")
        self.labels()
        if self.fw.get("element_dimensions", True):
            self.element_dimensions()
        if self.fw.get("internal_dimensions", True):
            self.internal_dimensions(quiet_columns=self.fw.get("element_dimensions", True))
        self.axes_and_dimensions()
        self.legend_and_notes(self.fw.get("title", "FORMWORK PLAN"))


PLAN_TITLES = {("B", "x"): "BOTTOM REINFORCEMENT - DIRECTION X", ("B", "y"): "BOTTOM REINFORCEMENT - DIRECTION Y",
               ("T", "x"): "TOP REINFORCEMENT - DIRECTION X", ("T", "y"): "TOP REINFORCEMENT - DIRECTION Y"}


def plan_wanted(g, plan):
    face, dirn = plan
    if g.kind == "perim":
        return face == "T"
    if g.kind == "ubar":
        return face == "B" and g.dirn == dirn
    return g.face == face and g.dirn == dirn


def draw_plan(doc, groups, chains, cfg, P, G, laps, plan, off):
    dc = cfg.get("drawing", {})
    fwp = Formwork(doc, cfg, P, G, off)
    if dc.get("hatch_on_plans", True):
        fwp.hatch_elements()                 # walls and columns hatched as on the formwork plan
    dr = Drawer(doc, cfg, P, G, off)
    if dr.show_range:
        dr.plan_ranges([g for g in groups if plan_wanted(g, plan)])
    for ch in chains:
        if (ch["face"], ch["dirn"]) == plan:
            dr.draw_chain(ch)
    for g in groups:
        if not plan_wanted(g, plan):
            continue
        if g.kind == "bar":
            dr.draw_bar_group(g)
        elif g.kind == "ubar":
            dr.draw_ubar_group(g)
        elif g.kind == "perim":
            dr.draw_perim(g)
    # the dimensions of the formwork plan, on every reinforcement plan too
    if dc.get("element_dimensions_on_plans", True):
        fwp.element_dimensions()
    if dc.get("internal_dimensions_on_plans", True):
        fwp.internal_dimensions(quiet_columns=dc.get("element_dimensions_on_plans", True))
    if dc.get("axes_on_plans", True):
        # the same axes and bubbles as on the formwork plan (with its outer dimension chains)
        fwp.axes_and_dimensions(with_dims=bool(dc.get("axis_dimensions_on_plans", True)))
    dr.notes(P, laps, face=plan[0], title=PLAN_TITLES[plan])


CLEAN = {}


def clean_source(doc, G, cfg):
    """If the slab came from an FE mesh or the supports from points, take those out of the output
    drawing and draw clean outlines instead (slab, openings, walls, columns)."""
    if not (G.from_mesh or G.from_points):
        return
    L = cfg["layers"]

    def lay(k):
        v = L.get(k) or []
        return {x.upper() for x in ([v] if isinstance(v, str) else v)}
    msp = doc.modelspace()
    kill = []
    for e in msp:
        lname = e.dxf.layer.upper()
        if G.from_mesh and lname in lay("slab") and e.dxftype() in ("3DFACE", "SOLID", "TRACE"):
            kill.append(e)
        elif G.from_points and lname in lay("support_points"):
            kill.append(e)
    for e in kill:
        msp.delete_entity(e)
    for name, col in (("SLAB", 3), ("OPENING", 1), ("WALL", 1), ("COLUMN", 4)):
        if name not in doc.layers:
            doc.layers.add(name, color=col)

    def poly(pg, layer):
        msp.add_lwpolyline([(x / G.u, y / G.u) for x, y in pg.exterior.coords], close=True,
                           dxfattribs={"layer": layer})
    for g in getattr(G.slab, "geoms", [G.slab]):
        poly(g, "SLAB")
        for r in g.interiors:
            poly(Polygon(r), "OPENING")
    if G.from_points:
        for g in getattr(G.walls_union, "geoms", [G.walls_union]):
            if not g.is_empty:
                poly(g, "WALL")
        for c in G.columns:
            poly(c.poly, "COLUMN")


def new_doc(src):
    doc = ezdxf.readfile(src)
    if CLEAN.get("G") is not None:
        clean_source(doc, CLEAN["G"], CLEAN["cfg"])
    if "DASHED" not in doc.linetypes:
        try:
            doc.linetypes.add("DASHED", pattern=[0.5, 0.25, -0.25], description="Dashed __ __")
        except Exception:
            pass
    return doc


def copy_geometry(doc, ents, off_units):
    """Copy the original drawing (slab, columns, axes, ...) to another plan position."""
    msp = doc.modelspace()
    failed = 0
    for e in ents:
        try:
            c = e.copy()
            c.translate(off_units[0], off_units[1], 0)
            msp.add_entity(c)
        except Exception:
            failed += 1
    if failed:
        warn(f"{failed} drawing object(s) (e.g. dimensions) could not be copied to the other plans")


PLAN_ORDER = [("B", "x"), ("B", "y"), ("T", "x"), ("T", "y")]


def plan_offsets(G, P, cfg, doc=None):
    """Layout:  formwork plan (top row), bottom X | bottom Y, top X | top Y.
    Gaps are taken from the extents of the whole drawing (axes, dimensions...)."""
    x0, y0, x1, y1 = G.bounds
    th = float(cfg.get("drawing", {}).get("text_height", 150))
    if doc is not None:
        try:
            from ezdxf import bbox
            ext = bbox.extents(doc.modelspace(), fast=True)
            if ext.has_data:
                x0, y0 = min(x0, ext.extmin.x * G.u), min(y0, ext.extmin.y * G.u)
                x1, y1 = max(x1, ext.extmax.x * G.u), max(y1, ext.extmax.y * G.u)
        except Exception:
            pass
    W, H = x1 - x0, y1 - y0
    gx = max(0.12 * W, 30 * th) + 60 * th          # room for dimensions, bubbles, legend
    gy = max(0.12 * H, 30 * th) + 40 * th
    offs = {("B", "x"): (0.0, 0.0), ("B", "y"): (W + gx, 0.0),
            ("T", "x"): (0.0, -(H + gy)), ("T", "y"): (W + gx, -(H + gy))}
    offs["FORMWORK"] = (0.0, H + gy)
    return offs


def formwork_enabled(cfg):
    return bool((cfg.get("formwork") or {}).get("enabled", True))


def write_plans(src, out, groups, chains, cfg, P, G, laps):
    """One DXF containing the formwork plan and the four reinforcement plans."""
    doc = new_doc(src)
    ents = list(doc.modelspace())
    offs = plan_offsets(G, P, cfg, doc)
    if formwork_enabled(cfg):
        off = offs["FORMWORK"]
        copy_geometry(doc, ents, (off[0] / G.u, off[1] / G.u))
        Formwork(doc, cfg, P, G, off).draw()
    for plan in PLAN_ORDER:
        off = offs[plan]
        if off != (0.0, 0.0):
            copy_geometry(doc, ents, (off[0] / G.u, off[1] / G.u))
        draw_plan(doc, groups, chains, cfg, P, G, laps, plan, off)
    doc.saveas(out)


def write_single_plan(src, out, groups, chains, cfg, P, G, laps, plan):
    doc = new_doc(src)
    draw_plan(doc, groups, chains, cfg, P, G, laps, plan, (0.0, 0.0))
    doc.saveas(out)


# =============================================================================
# Schedule, report, checks
# =============================================================================
def kg_per_m(d):
    return math.pi * d * d / 4 * 7.85e-3


FAM_NAMES = {"MESH": "Mesh", "BAND": "Grid band", "DIST": "Distribution (top)", "COL": "Over column", "WALL": "Over wall",
             "INTEG": "Integrity", "TRIM": "Opening trimmer", "UBAR": "Free-edge U-bar"}


def write_schedule(path, groups, sym):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Bar schedule"
    hdr = ["Mark", "Family", "Face", "Dir.", "Ø (mm)", "Shape", "Length (mm)", "No. of bars",
           "Total length (m)", "kg/m", "Weight (kg)", "Label on drawing"]
    ws.append(hdr)
    rows = [g for g in groups if g.kind in ("bar", "ubar")]
    rows.sort(key=lambda g: (int(g.mark[1:]), g.face, g.dirn))
    for g in rows:
        r = ws.max_row + 1
        fam = {"COL": "Under column", "WALL": "Under wall"}.get(g.fam) if is_mat() else None
        ws.append([g.mark, fam or FAM_NAMES.get(g.fam, g.fam), "Top" if g.face == "T" else "Bottom",
                   g.dirn.upper() if g.kind == "bar" else "edge", g.dia,
                   g.shape + (" (variable, avg. length)" if g.geom.get("var") else ""), g.length, g.n,
                   f"=G{r}*H{r}/1000", round(kg_per_m(g.dia), 3), f"=I{r}*J{r}",
                   (f"{g.n}{sym}{g.dia}x({g.geom['var'][0] / 10:.0f}÷{g.geom['var'][1] / 10:.0f})/{g.step / 10:.0f}"
                    if g.geom.get("var") else f"{g.n}{sym}{g.dia}x{g.length / 10:.0f}/{g.step / 10:.0f}")])
    last = ws.max_row
    ws.append([])
    ws.append(["TOTAL", "", "", "", "", "", "", f"=SUM(H2:H{last})", f"=SUM(I2:I{last})", "",
               f"=SUM(K2:K{last})"])
    bold = Font(bold=True)
    fill = PatternFill("solid", start_color="DDE7F0")
    for c in ws[1]:
        c.font, c.fill = bold, fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for c in ws[ws.max_row]:
        c.font = bold
    for col, w in zip("ABCDEFGHIJKL", [7, 16, 8, 6, 7, 9, 12, 10, 14, 8, 12, 22]):
        ws.column_dimensions[col].width = w
    for r in range(2, last + 1):
        ws[f"I{r}"].number_format = "0.00"
        ws[f"K{r}"].number_format = "0.0"

    ws2 = wb.create_sheet("Summary by Ø")
    ws2.append(["Ø (mm)", "Total length (m)", "kg/m", "Weight (kg)"])
    for c in ws2[1]:
        c.font, c.fill = bold, fill
    for d in sorted({g.dia for g in rows}):
        r = ws2.max_row + 1
        ws2.append([d, f"=SUMIF('Bar schedule'!E2:E{last},A{r},'Bar schedule'!I2:I{last})",
                    round(kg_per_m(d), 3), f"=B{r}*C{r}"])
        ws2[f"B{r}"].number_format = "0.0"
        ws2[f"D{r}"].number_format = "0.0"
    r = ws2.max_row + 1
    ws2.append(["TOTAL", f"=SUM(B2:B{r - 1})", "", f"=SUM(D2:D{r - 1})"])
    for c in ws2[r]:
        c.font = bold
    ws2[f"B{r}"].number_format = "0.0"
    ws2[f"D{r}"].number_format = "0.0"
    for col, w in zip("ABCD", [9, 18, 8, 14]):
        ws2.column_dimensions[col].width = w
    wb.save(path)


def checks(cfg, P):
    out = []
    asmin_ratio = max(0.26 * P.fctm / P.fyk, 0.0013)
    asmin = asmin_ratio * 1000 * P.d
    out.append(f"As,min (EC2 9.2.1.1) = max(0.26 fctm/fyk, 0.0013) b d = {asmin:.0f} mm2/m  (d = {P.d:.0f} mm)")
    s_main = min(3 * P.h, 400)
    s_peak = min(2 * P.h, 250)
    out.append(f"Max spacing (EC2 9.3.1.1(3)): {s_main:.0f} mm generally, {s_peak:.0f} mm in zones of max moment")

    def chk(name, spec, peak=False):
        if not spec:
            return
        d, s = spec["dia"], spec["spacing"]
        As = math.pi * d * d / 4 * 1000 / s
        lim = s_peak if peak else s_main
        ok_as = "OK" if As >= asmin else "BELOW As,min"
        ok_s = "OK" if s <= lim else f"SPACING > {lim:.0f} mm"
        out.append(f"  {name:<26} ф{d}/{s:.0f}: As = {As:.0f} mm2/m  -> {ok_as}; spacing {ok_s}")
        if ok_as != "OK" or ok_s != "OK":
            warn(f"{name} ф{d}/{s:.0f}: {ok_as if ok_as != 'OK' else ''} {ok_s if ok_s != 'OK' else ''}".strip())

    bm = cfg.get("bottom_mesh") or {}
    chk("Bottom mesh X", bm.get("x"))
    chk("Bottom mesh Y", bm.get("y"))
    lay = top_layout(cfg)
    if lay == "mesh":
        tm = cfg.get("top_mesh") or {}
        chk("Top mesh X", tm.get("x"))
        chk("Top mesh Y", tm.get("y"))
    else:
        gb = cfg.get("top_grid_bands") or {}
        chk("Top grid band X", gb.get("x"))
        chk("Top grid band Y", gb.get("y"))
        if lay == "bands_and_distribution":
            chk("Top distribution bars", cfg.get("top_distribution_bars"))
    cb = cfg.get("top_column_bars") or {}
    for k in ("internal", "edge", "corner"):
        if cb.get(k):
            s_eff = cb[k]["spacing"] / (2 if top_base(cfg, "x") else 1)
            out.append(f"  {'Under' if is_mat() else 'Over'} {k} columns: ф{cb[k]['dia']}/{cb[k]['spacing']:.0f} "
                       f"(combined spacing with band/mesh ≈ {s_eff:.0f} mm)")
            if s_eff > s_peak:
                warn(f"combined top spacing over {k} columns {s_eff:.0f} mm > {s_peak:.0f} mm")
    return out


def write_report(path, P, G, groups, laps, chk_lines, outputs):
    kinds = defaultdict(int)
    for c in G.columns:
        kinds[c.kind] += 1
    tot = defaultdict(float)
    for g in groups:
        if g.kind in ("bar", "ubar"):
            tot[g.dia] += g.n * g.length / 1000
    with open(path, "w", encoding="utf-8") as f:
        f.write(("MAT FOUNDATION" if is_mat() else "FLAT SLAB") + " DETAILING REPORT (EN 1992-1-1)\n" +
                "=" * 60 + "\n\n")
        if is_mat():
            f.write("Mat foundation: top bars lapped over the supports, bottom bars at mid-span.\n\n")
        f.write(f"Slab area: {G.slab.area / 1e6:.1f} m2,  openings: {len(G.holes)},  "
                f"walls (straight pieces): {len(G.walls)}\n")
        f.write(f"Columns: {len(G.columns)}  -> " + ", ".join(f"{k}: {v}" for k, v in kinds.items()) + "\n")
        f.write(f"Support lines X: {', '.join(f'{v:.0f}' for v in G.sup_lines['x'])}\n")
        f.write(f"Support lines Y: {', '.join(f'{v:.0f}' for v in G.sup_lines['y'])}\n\n")
        f.write(f"Materials: C{P.fck:.0f}, fctm={P.fctm:.2f}, fctd={P.fctd:.2f} MPa;  B{P.fyk:.0f}, fyd={P.fyd:.0f} MPa\n")
        f.write(f"h={P.h:.0f}, c_bot={P.cb:.0f}, c_top={P.ct:.0f}, d={P.d:.0f} mm\n\n")
        f.write("Anchorage / lap lengths (EC2 8.4, 8.7; sigma_sd = fyd, alpha1..5 = 1)\n")
        f.write("  Ø    face   eta1   fbd     lb,rqd   l_bd    l0\n")
        for d in sorted(laps):
            for face in ("B", "T"):
                e1, fbd, lbr = P.bond(d, face)
                f.write(f"  {d:<4} {'bottom' if face == 'B' else 'top':<6} {e1:<6.1f} {fbd:<7.2f} {lbr:<8.0f} "
                        f"{P.lbd(d, face):<7.0f} {P.lap(d, face):.0f}\n")
        f.write("\nChecks\n")
        for ln in chk_lines:
            f.write("  " + ln + "\n")
        f.write("\nSteel quantities\n")
        tw = 0
        for d in sorted(tot):
            w = tot[d] * kg_per_m(d)
            tw += w
            f.write(f"  ф{d:<3} {tot[d]:9.1f} m  {w:9.1f} kg\n")
        f.write(f"  TOTAL {tw:13.1f} kg   ({tw / (G.slab.area / 1e6):.1f} kg/m2)\n\n")
        f.write("Warnings\n")
        for w in WARNINGS or ["none"]:
            f.write("  - " + w + "\n")
        f.write("\nFiles\n")
        for o in outputs:
            f.write("  " + o + "\n")


def preview(dxf_path, png_path):
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        from ezdxf.addons.drawing import Frontend, RenderContext
        from ezdxf.addons.drawing.matplotlib import MatplotlibBackend
        from ezdxf.addons.drawing.config import (BackgroundPolicy, ColorPolicy, Configuration,
                                                 LineweightPolicy)
        doc = ezdxf.readfile(dxf_path)
        fig = plt.figure(figsize=(24, 24))
        ax = fig.add_axes([0, 0, 1, 1])
        ctx = RenderContext(doc)
        cfgd = Configuration(lineweight_policy=LineweightPolicy.ABSOLUTE, lineweight_scaling=1.0,
                             background_policy=BackgroundPolicy.WHITE)
        Frontend(ctx, MatplotlibBackend(ax), config=cfgd).draw_layout(doc.modelspace(), finalize=True)
        fig.savefig(png_path, dpi=150, facecolor="white")
        plt.close(fig)
        return True
    except Exception as ex:
        warn(f"preview failed: {ex}")
        return False


# =============================================================================
# Main
# =============================================================================
def run(dxf, cfg_path, outdir, structure=None):
    print(f"Reading config {cfg_path}")
    with open(cfg_path, encoding="utf-8") as f:
        cfg = yaml.safe_load(f)
    if structure:
        cfg["structure"] = structure
    STRUCTURE[0] = cfg.get("structure", "flat_slab")
    if STRUCTURE[0] not in ("flat_slab", "mat_foundation"):
        sys.exit(f"ERROR: structure must be flat_slab or mat_foundation, not {STRUCTURE[0]!r}")
    mat = is_mat()
    if mat:
        print("  MAT FOUNDATION: top bars lapped over the supports, bottom bars at mid-span")
        if top_layout(cfg) != "mesh":
            warn("mat foundation: only top_layout: mesh is used (top and bottom meshes)")
        cfg["top_layout"] = "mesh"
        if (cfg.get("integrity_bars") or {}).get("enabled"):
            warn("mat foundation: integrity bars (EC2 9.4.1(3)) are for flat slabs - left out")
            cfg["integrity_bars"]["enabled"] = False
    P = Params(cfg)
    print(f"Reading geometry {dxf}")
    src_doc = ezdxf.readfile(dxf)
    G = Geometry(src_doc, cfg, P)
    CLEAN["G"], CLEAN["cfg"] = G, cfg
    print(f"  slab {G.slab.area / 1e6:.1f} m2, {len(G.columns)} columns, {len(G.walls)} wall pieces, "
          f"{len(G.holes)} openings")
    prefix = cfg.get("drawing", {}).get("layer_prefix", "REB_")

    # laps: flat slab - bottom over the supports, top at mid-span; mat foundation - the other way round
    bars: list[Bar] = []
    for dirn in ("x", "y"):
        spec = (cfg.get("bottom_mesh") or {}).get(dirn)
        if spec:
            bars += gen_mesh(G, P, spec, "B", dirn, "MESH", G.sup_lines[dirn], mat)
        spec = (cfg.get("top_mesh") or {}).get(dirn)
        if spec and top_layout(cfg) == "mesh":
            bars += gen_mesh(G, P, spec, "T", dirn, "MESH", G.sup_lines[dirn], not mat)
    bars += gen_grid_bands(G, P, cfg)
    bars += gen_top_distribution(G, P, cfg, bars)
    link_top_chains(bars)
    support = gen_column_bars(G, P, cfg) + gen_wall_bars(G, P, cfg)
    if mat:                                 # under the columns / walls of a mat: bottom bars
        for b in support:
            b.face = "B"
    bars += support
    bars += gen_integrity(G, P, cfg)
    bars += gen_trimmers(G, P, cfg)
    add_hooks(bars, G, P)
    for b in bars:
        if b.length > P.stock + 1:
            warn(f"a {b.fam} bar ф{b.dia} is {b.length:.0f} mm long (> stock {P.stock:.0f} mm)")

    MERGE_MAX[0] = int(cfg.get("merge_small_groups", 6))
    groups = merge_variable(group_bars(bars, prefix), prefix)
    chains = build_chains(bars, groups)
    groups += gen_ubars(G, P, cfg, prefix)
    groups += gen_perimeters(G, P, cfg, prefix)
    assign_marks(groups)
    # draw big families first so their labels get the best spots
    order = {"MESH": 0, "INTEG": 1, "BAND": 2, "DIST": 2, "COL": 3, "WALL": 4, "TRIM": 5, "UBAR": 6, "PUNCH": 7}
    groups.sort(key=lambda g: (order.get(g.fam, 9), -g.n))

    dias = sorted({g.dia for g in groups if g.kind in ("bar", "ubar")})
    laps = {d: (P.lap(d, "B"), P.lap(d, "T")) for d in dias}
    chk_lines = checks(cfg, P)

    os.makedirs(outdir, exist_ok=True)
    base = os.path.splitext(os.path.basename(dxf))[0]
    if mat:
        base += "_MAT"
    out_plans = os.path.join(outdir, f"{base}_REINFORCEMENT_PLANS.dxf")
    out_xls = os.path.join(outdir, f"{base}_bar_schedule.xlsx")
    out_rep = os.path.join(outdir, f"{base}_report.txt")
    write_plans(dxf, out_plans, groups, chains, cfg, P, G, laps)
    outputs = [out_plans]
    if cfg.get("drawing", {}).get("separate_plan_files", False):
        names = {("B", "x"): "1_BOTTOM_X", ("B", "y"): "2_BOTTOM_Y", ("T", "x"): "3_TOP_X", ("T", "y"): "4_TOP_Y"}
        for plan in PLAN_ORDER:
            p_ = os.path.join(outdir, f"{base}_{names[plan]}.dxf")
            write_single_plan(dxf, p_, groups, chains, cfg, P, G, laps, plan)
            outputs.append(p_)
    write_schedule(out_xls, groups, cfg.get("drawing", {}).get("dia_symbol", "ф"))
    outputs += [out_xls, out_rep]
    if cfg.get("drawing", {}).get("preview_png", True):
        print("  rendering preview ...")
        png = out_plans[:-4] + ".png"
        if preview(out_plans, png):
            outputs.append(png)
    write_report(out_rep, P, G, groups, laps, chk_lines, outputs)
    nb = sum(g.n for g in groups if g.kind in ("bar", "ubar"))
    print(f"  {len(groups)} bar groups, {nb} bars, {len({g.mark for g in groups if g.mark})} bar marks")
    print("Written:")
    for o in outputs:
        print("  ", o)
    return outputs


def main():
    ap = argparse.ArgumentParser(description="Flat slab reinforcement detailer (EC2)")
    ap.add_argument("dxf", help="DXF with slab outline, openings, columns and walls")
    ap.add_argument("-c", "--config", default=None, help="config YAML (default: config.yaml next to this script)")
    ap.add_argument("-o", "--out", default="output", help="output folder")
    ap.add_argument("--list-layers", action="store_true",
                    help="only list the layers of the DXF and what is on them (to fill in the config)")
    a = ap.parse_args()
    if a.list_layers:
        doc = ezdxf.readfile(a.dxf)
        counts = defaultdict(lambda: defaultdict(int))
        for e, lay in iter_entities(doc.modelspace()):
            counts[lay][e.dxftype()] += 1
        units = {0: "unitless", 1: "inches", 4: "mm", 5: "cm", 6: "m"}.get(doc.header.get("$INSUNITS", 0), "?")
        print(f"Drawing units ($INSUNITS): {units}")
        for lay in sorted(counts):
            print(f"  {lay:<30} " + ", ".join(f"{k}:{v}" for k, v in sorted(counts[lay].items())))
        return
    cfg = a.config or os.path.join(os.path.dirname(os.path.abspath(__file__)), "config.yaml")
    run(a.dxf, cfg, a.out)


if __name__ == "__main__":
    main()
